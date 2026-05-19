import csv
import io
import logging
from datetime import timedelta

from django.db.models import Case, Count, DecimalField, F, Sum, When

logger = logging.getLogger(__name__)

LEAVE_CATEGORY = "HOLIDAYS/LEAVES"

# ---------------------------------------------------------------------------
# Registry
# ---------------------------------------------------------------------------

REPORT_REGISTRY = {}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _sprint_working_days(sprint):
    """Count Mon–Fri days within a sprint's start/end date range."""
    count = 0
    cur = sprint.start_date
    while cur <= sprint.end_date:
        if cur.weekday() < 5:
            count += 1
        cur += timedelta(days=1)
    return max(count, 1)


def _r2(v):
    return round(float(v or 0), 2)


# ---------------------------------------------------------------------------
# Report list service
# ---------------------------------------------------------------------------

class ReportService:
    @staticmethod
    def list_reports(report_type=None):
        from .models import Report

        qs = Report.objects.filter(is_active=True)
        if report_type:
            types = [t.strip().upper() for t in report_type.split(",")]
            qs = qs.filter(report_type__in=types)
        return list(qs)

    @staticmethod
    def get_report_by_slug(slug):
        from .models import Report

        return Report.objects.get(slug=slug, is_active=True)

    @staticmethod
    def create_custom_report(data, user=None):
        from .models import Report

        return Report.objects.create(
            slug=data["slug"],
            name=data["name"],
            description=data.get("description", ""),
            report_type=Report.CUSTOM,
            is_active=True,
            sort_order=data.get("sort_order", 0),
            created_by=user,
            updated_by=user,
        )


# ---------------------------------------------------------------------------
# Demand & Capacity — config service
# ---------------------------------------------------------------------------

class DemandCapacityConfigService:
    @staticmethod
    def get_or_create(plan_id, version_id, user=None):
        from apps.resource_plans.models import ResourcePlan, ResourcePlanVersion

        from .models import DemandCapacityConfig

        plan = ResourcePlan.objects.get(pk=plan_id)
        version = ResourcePlanVersion.objects.get(pk=version_id)
        config, created = DemandCapacityConfig.objects.get_or_create(
            plan=plan,
            version=version,
            defaults={"created_by": user, "updated_by": user},
        )
        if not created and user:
            config.updated_by = user
            config.save(update_fields=["updated_by", "updated_at"])
        return config

    @staticmethod
    def get_configure_data(plan_id, version_id):
        from apps.programmes.models import Programme
        from apps.resource_plans.models import (
            ResourcePlan,
            ResourcePlanVersion,
            ResourcePlanVersionProject,
        )

        from .models import DemandCapacityConfig

        try:
            plan = ResourcePlan.objects.get(pk=plan_id)
            version = ResourcePlanVersion.objects.get(pk=version_id)
        except (ResourcePlan.DoesNotExist, ResourcePlanVersion.DoesNotExist):
            return None

        programme_ids = set(
            ResourcePlanVersionProject.objects.filter(
                version=version,
                project__programme__isnull=False,
            ).values_list("project__programme_id", flat=True)
        )
        programmes = list(Programme.objects.filter(id__in=programme_ids).order_by("name"))

        try:
            config = DemandCapacityConfig.objects.get(plan=plan, version=version)
            mappings = list(config.mappings.select_related("programme").all())
            mapping_dict = {m.programme_id: m for m in mappings}
        except DemandCapacityConfig.DoesNotExist:
            config = None
            mapping_dict = {}

        return {
            "plan": plan,
            "version": version,
            "config": config,
            "programmes": programmes,
            "mapping_dict": mapping_dict,
        }

    @staticmethod
    def list_mappings(plan_id, version_id):
        from .models import ProgrammeCategoryMapping

        return (
            ProgrammeCategoryMapping.objects.filter(
                config__plan_id=plan_id,
                config__version_id=version_id,
            )
            .select_related("programme", "config")
            .order_by("programme__name")
        )

    @staticmethod
    def upsert_mapping(plan_id, version_id, programme_id, category_label, user=None):
        from apps.programmes.models import Programme

        from .models import ProgrammeCategoryMapping

        config = DemandCapacityConfigService.get_or_create(plan_id, version_id, user)
        programme = Programme.objects.get(pk=programme_id)
        mapping, created = ProgrammeCategoryMapping.objects.update_or_create(
            config=config,
            programme=programme,
            defaults={"category_label": category_label, "created_by": user},
        )
        return mapping, created

    @staticmethod
    def delete_mapping(mapping_id):
        from .models import ProgrammeCategoryMapping

        ProgrammeCategoryMapping.objects.filter(pk=mapping_id).delete()

    @staticmethod
    def bulk_upsert_mappings(plan_id, version_id, mappings, user=None):
        """
        Each item: {programme_id, category_label}.
        Empty label → delete existing mapping for that programme.
        """
        from .models import ProgrammeCategoryMapping

        config = DemandCapacityConfigService.get_or_create(plan_id, version_id, user)
        results = []
        for item in mappings:
            prog_id = item["programme_id"]
            label = item.get("category_label", "").strip()
            if not label:
                ProgrammeCategoryMapping.objects.filter(
                    config=config, programme_id=prog_id
                ).delete()
                continue
            mapping, created = ProgrammeCategoryMapping.objects.update_or_create(
                config=config,
                programme_id=prog_id,
                defaults={"category_label": label, "created_by": user},
            )
            results.append({"mapping": mapping, "created": created})
        return results


# ---------------------------------------------------------------------------
# Demand & Capacity — report data service
# ---------------------------------------------------------------------------

class DemandCapacityService:

    @staticmethod
    def get_data(plan_id, version_id, team_id=None, employee_type_id=None):
        from apps.delivery_teams.models import DeliveryTeam
        from apps.programmes.models import Programme
        from apps.resource_plans.models import (
            ResourcePlan,
            ResourcePlanAllocation,
            ResourcePlanAllocationSet,
            ResourcePlanMemberCapacity,
            ResourcePlanVersion,
        )
        from apps.sprints.models import Sprint

        from .models import DemandCapacityConfig

        try:
            plan = ResourcePlan.objects.select_related("financial_year").get(pk=plan_id)
            version = ResourcePlanVersion.objects.get(pk=version_id)
        except (ResourcePlan.DoesNotExist, ResourcePlanVersion.DoesNotExist):
            return None

        # ── Sprints & month structure ──────────────────────────────────────
        sprints = list(
            Sprint.objects.filter(financial_year=plan.financial_year)
            .order_by("sprint_number")
        )
        sprint_ids = [s.id for s in sprints]

        months_order = []          # ordered unique month keys
        sprint_to_month = {}       # sprint.id → month_key
        month_working_days = {}    # month_key → sum of business days across sprints
        month_labels = {}          # month_key → display label
        seen_months: set = set()

        for s in sprints:
            wd = _sprint_working_days(s)
            mkey = f"{s.end_date.year}-{s.month}"
            if mkey not in seen_months:
                seen_months.add(mkey)
                months_order.append(mkey)
                month_labels[mkey] = f"{s.month} '{str(s.end_date.year)[2:]}"
                month_working_days[mkey] = 0
            sprint_to_month[s.id] = mkey
            month_working_days[mkey] += wd

        # ── Category mappings (fresh from DB — no cache) ───────────────────
        try:
            config = DemandCapacityConfig.objects.get(plan=plan, version=version)
            mapping_dict = {
                m.programme_id: m.category_label
                for m in config.mappings.all()
            }
        except DemandCapacityConfig.DoesNotExist:
            mapping_dict = {}

        def _category(prog_id, prog_names):
            return mapping_dict.get(prog_id, prog_names.get(prog_id, str(prog_id)))

        # ── Active allocation set ──────────────────────────────────────────
        alloc_set = (
            ResourcePlanAllocationSet.objects.filter(
                version=version,
                status=ResourcePlanAllocationSet.STATUS_ACTIVE,
            )
            .order_by("-created_at")
            .first()
        )

        # ── Demand: raw days per (team, sprint, programme) ─────────────────
        demand_rows = []
        if alloc_set:
            qs = ResourcePlanAllocation.objects.filter(
                allocation_set=alloc_set,
                programme__isnull=False,
                sprint_id__in=sprint_ids,
                placeholder_engineer__isnull=True,  # Exclude AUTO #X placeholder slots
            )
            if team_id:
                qs = qs.filter(team_id=team_id)
            if employee_type_id:
                qs = qs.filter(team_member__employment_type_id=employee_type_id)

            demand_rows = list(
                qs.values("team_id", "sprint_id", "programme_id").annotate(
                    total=Sum(
                        Case(
                            When(override_days__isnull=False, then=F("override_days")),
                            default=F("engine_days"),
                            output_field=DecimalField(),
                        )
                    )
                )
            )

        prog_ids_needed = {r["programme_id"] for r in demand_rows}
        prog_names = {
            p.id: p.name
            for p in Programme.objects.filter(id__in=prog_ids_needed)
        }

        # ── Capacity & leaves: gross working_days per (team, sprint) ──────
        cap_qs = ResourcePlanMemberCapacity.objects.filter(
            version=version, sprint_id__in=sprint_ids
        )
        if team_id:
            cap_qs = cap_qs.filter(team_member__team_assignments__team_id=team_id)
        if employee_type_id:
            cap_qs = cap_qs.filter(team_member__employment_type_id=employee_type_id)

        cap_rows = list(
            cap_qs.values("sprint_id", "team_member__team_assignments__team_id").annotate(
                working=Sum("working_days"),
                leaves=Sum(F("leave_days") + F("holiday_days") + F("placeholder_days")),
            )
        )

        # ── Aggregate demand into: {team_id: {category: {month_key: days}}} ─
        def _build_cat_days(rows, team_filter=None):
            """Accumulate raw demand days by category × month."""
            cat_days: dict = {}
            for r in rows:
                if team_filter is not None and r["team_id"] != team_filter:
                    continue
                mkey = sprint_to_month.get(r["sprint_id"])
                if not mkey:
                    continue
                cat = _category(r["programme_id"], prog_names)
                cat_days.setdefault(cat, {})
                cat_days[cat][mkey] = cat_days[cat].get(mkey, 0.0) + float(r["total"] or 0)
            return cat_days

        def _days_to_fte(cat_days):
            """Convert {cat: {month: days}} → {cat: {month: fte}}."""
            out = {}
            for cat, md in cat_days.items():
                out[cat] = {}
                for mkey, days in md.items():
                    wd = month_working_days.get(mkey, 1)
                    out[cat][mkey] = _r2(days / wd) if wd else 0.0
            return out

        def _sprint_to_month_raw(by_sprint):
            """Sum raw values per month, then convert to FTE."""
            by_month: dict = {}
            for sid, val in by_sprint.items():
                mkey = sprint_to_month.get(sid)
                if mkey:
                    by_month[mkey] = by_month.get(mkey, 0.0) + val
            return {
                mkey: _r2(days / month_working_days.get(mkey, 1))
                for mkey, days in by_month.items()
            }

        # Overall demand FTE by category
        overall_cat_days = _build_cat_days(demand_rows)
        overall_demand_fte = _days_to_fte(overall_cat_days)

        # Seed all configured category labels so they appear even with 0 demand
        for cat_label in mapping_dict.values():
            if cat_label and cat_label not in overall_demand_fte:
                overall_demand_fte[cat_label] = {}

        # Overall capacity + leave FTE
        cap_by_sprint: dict = {}
        leave_by_sprint: dict = {}
        team_cap_by_sprint: dict = {}   # {(tid, sid): working_days}
        team_leave_by_sprint: dict = {} # {(tid, sid): leave_days}

        for r in cap_rows:
            sid = r["sprint_id"]
            tid = r["team_member__team_assignments__team_id"]
            w = float(r["working"] or 0)
            lv = float(r["leaves"] or 0)
            cap_by_sprint[sid] = cap_by_sprint.get(sid, 0.0) + w
            leave_by_sprint[sid] = leave_by_sprint.get(sid, 0.0) + lv
            team_cap_by_sprint[(tid, sid)] = team_cap_by_sprint.get((tid, sid), 0.0) + w
            team_leave_by_sprint[(tid, sid)] = team_leave_by_sprint.get((tid, sid), 0.0) + lv

        overall_cap_fte = _sprint_to_month_raw(cap_by_sprint)
        overall_leave_fte = _sprint_to_month_raw(leave_by_sprint)

        # Add HOLIDAYS/LEAVES as a demand-side category so it stacks in the chart
        if overall_leave_fte:
            overall_demand_fte[LEAVE_CATEGORY] = overall_leave_fte

        # Build ordered categories: sorted regular + HOLIDAYS/LEAVES last
        regular_cats = sorted(
            cat for cat in overall_demand_fte if cat != LEAVE_CATEGORY
        )
        all_categories = regular_cats + (
            [LEAVE_CATEGORY] if LEAVE_CATEGORY in overall_demand_fte else []
        )

        # ── Per-team breakdown ─────────────────────────────────────────────
        team_ids = set()
        for r in demand_rows:
            if r["team_id"]:
                team_ids.add(r["team_id"])
        for (tid, _) in team_cap_by_sprint:
            if tid:
                team_ids.add(tid)

        teams_data = []
        if team_ids:
            teams_qs = DeliveryTeam.objects.filter(id__in=team_ids).order_by("name")
            for team in teams_qs:
                tid = team.id

                t_cat_days = _build_cat_days(demand_rows, team_filter=tid)
                t_demand_fte = _days_to_fte(t_cat_days)

                t_cap_raw = {
                    sid: v for (t, sid), v in team_cap_by_sprint.items() if t == tid
                }
                t_leave_raw = {
                    sid: v for (t, sid), v in team_leave_by_sprint.items() if t == tid
                }
                t_cap_fte = _sprint_to_month_raw(t_cap_raw)
                t_leave_fte = _sprint_to_month_raw(t_leave_raw)

                if t_leave_fte:
                    t_demand_fte[LEAVE_CATEGORY] = t_leave_fte

                t_demand_by_month = _totals_by_month(t_demand_fte, months_order)
                t_risk, t_util = _risk_util(t_demand_by_month, t_cap_fte, months_order)

                teams_data.append({
                    "id": tid,
                    "name": team.name,
                    "demand": t_demand_fte,
                    "capacity": t_cap_fte,
                    "totals": {
                        "demand_by_month": t_demand_by_month,
                        "demand_total": _r2(sum(t_demand_by_month.values())),
                        "capacity_total": _r2(sum(t_cap_fte.values())),
                        "risk_by_month": t_risk,
                        "utilisation_by_month": t_util,
                    },
                })

        # ── Overall totals ─────────────────────────────────────────────────
        demand_by_month = _totals_by_month(overall_demand_fte, months_order)
        risk_by_month, util_by_month = _risk_util(
            demand_by_month, overall_cap_fte, months_order
        )

        months_list = [
            {
                "key": k,
                "label": month_labels[k],
                "working_days": month_working_days[k],
            }
            for k in months_order
        ]

        return {
            "plan": {"id": plan.id, "name": plan.name},
            "version": {
                "id": version.id,
                "version": version.version,
                "status": version.status,
            },
            "months": months_list,
            "categories": all_categories,
            "demand": overall_demand_fte,
            "capacity": overall_cap_fte,
            "totals": {
                "demand_by_month": demand_by_month,
                "demand_total": _r2(sum(demand_by_month.values())),
                "capacity_total": _r2(sum(overall_cap_fte.values())),
                "risk_by_month": risk_by_month,
                "utilisation_by_month": util_by_month,
            },
            "teams": teams_data,
            "has_allocation_set": alloc_set is not None,
        }

    # ── Export helpers ─────────────────────────────────────────────────────

    @staticmethod
    def export_csv(data: dict) -> bytes:
        months = data["months"]
        categories = data["categories"]
        demand = data["demand"]
        capacity = data["capacity"]
        totals = data["totals"]

        buf = io.StringIO()
        w = csv.writer(buf)

        # Header
        w.writerow(["Category"] + [m["label"] for m in months] + ["Total"])

        # Demand rows
        for cat in categories:
            cat_data = demand.get(cat, {})
            row_vals = [cat_data.get(m["key"], 0.0) for m in months]
            w.writerow([cat] + [f"{v:.2f}" for v in row_vals] + [f"{sum(row_vals):.2f}"])

        # Total Demand
        td = [totals["demand_by_month"].get(m["key"], 0.0) for m in months]
        w.writerow(["Total Demand"] + [f"{v:.2f}" for v in td] + [f"{sum(td):.2f}"])

        # Total Capacity
        tc = [capacity.get(m["key"], 0.0) for m in months]
        w.writerow(["Total Capacity"] + [f"{v:.2f}" for v in tc] + [f"{sum(tc):.2f}"])

        # FTE Risk
        tr = [totals["risk_by_month"].get(m["key"], 0.0) for m in months]
        w.writerow(["FTE Risk"] + [f"{v:.2f}" for v in tr] + [f"{sum(tr):.2f}"])

        # Utilisation %
        tu = [totals["utilisation_by_month"].get(m["key"], 0.0) for m in months]
        w.writerow(["Utilisation %"] + [f"{v:.2f}%" for v in tu] + [""])

        return buf.getvalue().encode("utf-8-sig")

    @staticmethod
    def export_xlsx(data: dict) -> bytes:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        months = data["months"]
        categories = data["categories"]
        demand = data["demand"]
        capacity = data["capacity"]
        totals = data["totals"]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Demand & Capacity"

        H_FILL  = PatternFill("solid", fgColor="1F4E79")
        TD_FILL = PatternFill("solid", fgColor="DBEAFE")
        TC_FILL = PatternFill("solid", fgColor="D1FAE5")
        RK_FILL = PatternFill("solid", fgColor="FEF3C7")
        LV_FILL = PatternFill("solid", fgColor="F3F4F6")
        CENTER  = Alignment(horizontal="center")

        def hdr_cell(row, col, val):
            c = ws.cell(row, col, val)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = H_FILL
            c.alignment = CENTER
            return c

        def styled_row(row, col_start, values, fill, bold=False, fmt="{:.2f}"):
            for i, v in enumerate(values):
                c = ws.cell(row, col_start + i, float(v) if isinstance(v, (int, float)) else v)
                c.fill = fill
                c.alignment = CENTER
                if bold:
                    c.font = Font(bold=True)

        n = len(months)
        total_col = n + 2

        # Header row
        hdr_cell(1, 1, "Category")
        for ci, m in enumerate(months, 2):
            hdr_cell(1, ci, m["label"])
        hdr_cell(1, total_col, "Total")

        # Category demand rows
        PALETTE = [
            "BFD7FF","C6F6D5","FDE68A","FCA5A5","DDD6FE",
            "A5F3FC","FED7AA","D9F99D","FBCFE8","C7D2FE",
        ]
        next_row = 2
        for idx, cat in enumerate(categories):
            cat_data = demand.get(cat, {})
            fill_hex = "F9FAFB" if cat == LEAVE_CATEGORY else PALETTE[idx % len(PALETTE)]
            fill = PatternFill("solid", fgColor=fill_hex)
            ws.cell(next_row, 1, cat).font = Font(italic=(cat == LEAVE_CATEGORY))
            ws.cell(next_row, 1).fill = fill
            row_vals = [cat_data.get(m["key"], 0.0) for m in months]
            for ci, v in enumerate(row_vals, 2):
                c = ws.cell(next_row, ci, round(v, 2))
                c.alignment = CENTER
                c.fill = fill
            ws.cell(next_row, total_col, round(sum(row_vals), 2)).alignment = CENTER
            ws.cell(next_row, total_col).fill = fill
            next_row += 1

        # Total Demand
        td = [totals["demand_by_month"].get(m["key"], 0.0) for m in months]
        ws.cell(next_row, 1, "Total Demand").font = Font(bold=True)
        ws.cell(next_row, 1).fill = TD_FILL
        styled_row(next_row, 2, td + [sum(td)], TD_FILL, bold=True)
        next_row += 1

        # Total Capacity
        tc = [capacity.get(m["key"], 0.0) for m in months]
        ws.cell(next_row, 1, "Total Capacity").font = Font(bold=True)
        ws.cell(next_row, 1).fill = TC_FILL
        styled_row(next_row, 2, tc + [sum(tc)], TC_FILL, bold=True)
        next_row += 1

        # FTE Risk
        tr = [totals["risk_by_month"].get(m["key"], 0.0) for m in months]
        ws.cell(next_row, 1, "FTE Risk").font = Font(bold=True)
        ws.cell(next_row, 1).fill = RK_FILL
        styled_row(next_row, 2, tr + [sum(tr)], RK_FILL, bold=True)
        next_row += 1

        # Utilisation %
        tu = [totals["utilisation_by_month"].get(m["key"], 0.0) for m in months]
        ws.cell(next_row, 1, "Utilisation %").font = Font(bold=True, italic=True)
        for ci, v in enumerate(tu, 2):
            c = ws.cell(next_row, ci, f"{v:.2f}%")
            c.alignment = CENTER
        ws.cell(next_row, total_col, "")

        # Column widths
        ws.column_dimensions["A"].width = 26
        for ci in range(2, total_col + 1):
            ws.column_dimensions[get_column_letter(ci)].width = 13

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()


# ---------------------------------------------------------------------------
# Private helpers used by get_data
# ---------------------------------------------------------------------------

def _totals_by_month(demand_fte: dict, months_order: list) -> dict:
    """Sum all category FTEs per month."""
    by_month: dict = {}
    for cat_data in demand_fte.values():
        for mkey, fte in cat_data.items():
            by_month[mkey] = by_month.get(mkey, 0.0) + fte
    return {k: _r2(by_month.get(k, 0.0)) for k in months_order}


def _risk_util(demand_by_month: dict, cap_fte: dict, months_order: list):
    """Return (risk_by_month, utilisation_by_month) dicts."""
    risk: dict = {}
    util: dict = {}
    for mkey in months_order:
        d = demand_by_month.get(mkey, 0.0)
        c = cap_fte.get(mkey, 0.0)
        risk[mkey] = _r2(c - d)
        util[mkey] = _r2(d / c * 100) if c else 0.0
    return risk, util


# ---------------------------------------------------------------------------
# Sprint Forecast vs. Actuals — report data service
# ---------------------------------------------------------------------------

class SprintForecastActualsService:

    @staticmethod
    def get_data(sprint_id, team_id=None):
        from decimal import Decimal

        from apps.sprint_forecast.models import (
            IMPORT_TYPE_FORECAST,
            IMPORT_TYPE_ACTUAL,
            SprintConfirmedRow,
            SprintImportReviewComplete,
        )
        from apps.sprints.models import Sprint

        try:
            sprint = Sprint.objects.get(pk=sprint_id)
        except Sprint.DoesNotExist:
            return None

        forecast_rc = SprintImportReviewComplete.objects.filter(
            sprint=sprint, import_type=IMPORT_TYPE_FORECAST
        ).first()
        actual_rc = SprintImportReviewComplete.objects.filter(
            sprint=sprint, import_type=IMPORT_TYPE_ACTUAL
        ).first()

        qs = SprintConfirmedRow.objects.filter(sprint=sprint).select_related(
            'team', 'assignee', 'label__project__programme', 'mapping'
        )
        if team_id:
            qs = qs.filter(team_id=team_id)

        fc_rows = [r for r in qs if r.import_type == IMPORT_TYPE_FORECAST]
        ac_rows = [r for r in qs if r.import_type == IMPORT_TYPE_ACTUAL]

        has_forecast = len(fc_rows) > 0
        has_actuals = len(ac_rows) > 0

        # Build lookup dicts keyed by (team_id, assignee_key, label_key, mapping_key)
        def _row_key(row):
            assignee_key = row.assignee_id if row.assignee_id else row.assignee_raw
            label_key = row.label_id if row.label_id else row.label_raw
            mapping_key = row.mapping_id if row.mapping_id else row.mapping_raw
            return (row.team_id, assignee_key, label_key, mapping_key)

        def _row_meta(row):
            label = row.label
            project = label.project if label else None
            programme = project.programme if project else None
            mapping = row.mapping
            return {
                'team_id': row.team_id,
                'team': row.team.name if row.team else '—',
                'assignee_id': row.assignee_id,
                'assignee': row.assignee.display_name if row.assignee else row.assignee_raw or '—',
                'label_id': row.label_id,
                'label': label.label if label else row.label_raw or '—',
                'project_id': project.id if project else None,
                'project': project.name if project else '—',
                'programme_id': programme.id if programme else None,
                'programme': programme.name if programme else '—',
                'mapping_id': row.mapping_id,
                'mapping': mapping.code if mapping else row.mapping_raw or '—',
                'mapping_name': mapping.name if mapping else row.mapping_raw or '—',
            }

        fc_by_key = {}
        for row in fc_rows:
            k = _row_key(row)
            fc_by_key.setdefault(k, {'meta': _row_meta(row), 'days': Decimal('0')})
            fc_by_key[k]['days'] += row.days

        ac_by_key = {}
        for row in ac_rows:
            k = _row_key(row)
            ac_by_key.setdefault(k, {'meta': _row_meta(row), 'days': Decimal('0')})
            ac_by_key[k]['days'] += row.days

        all_keys = set(fc_by_key) | set(ac_by_key)

        rows = []
        added = removed = changed = unchanged = 0
        forecast_total = Decimal('0')
        actual_total = Decimal('0')

        for k in all_keys:
            fc_entry = fc_by_key.get(k)
            ac_entry = ac_by_key.get(k)
            meta = (fc_entry or ac_entry)['meta']

            fc_days = fc_entry['days'] if fc_entry else Decimal('0')
            ac_days = ac_entry['days'] if ac_entry else Decimal('0')
            delta = ac_days - fc_days

            if not fc_entry:
                status = 'added'
                added += 1
            elif not ac_entry:
                status = 'removed'
                removed += 1
            elif fc_days != ac_days:
                status = 'changed'
                changed += 1
            else:
                status = 'unchanged'
                unchanged += 1

            forecast_total += fc_days
            actual_total += ac_days

            rows.append({
                **meta,
                'forecast_days': _r2(fc_days),
                'actual_days': _r2(ac_days),
                'delta': _r2(delta),
                'status': status,
            })

        rows.sort(key=lambda r: (r['team'], r['assignee'], r['label']))

        return {
            'sprint': {'id': sprint.id, 'name': sprint.sprint_name},
            'has_forecast': has_forecast,
            'has_actuals': has_actuals,
            'forecast_complete': forecast_rc is not None,
            'actual_complete': actual_rc is not None,
            'can_compare': has_forecast and has_actuals,
            'rows': rows,
            'summary': {
                'forecast_total_days': _r2(forecast_total),
                'actual_total_days': _r2(actual_total),
                'delta': _r2(actual_total - forecast_total),
                'added': added,
                'removed': removed,
                'changed': changed,
                'unchanged': unchanged,
            },
        }

    @staticmethod
    def export_csv(data: dict, sprint_name: str = '') -> bytes:
        rows = data.get('rows', [])
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([
            'Team', 'Engineer', 'Label', 'Project', 'Programme', 'Finance Type',
            'Forecast Days', 'Actual Days', 'Delta', 'Status',
        ])
        for r in rows:
            w.writerow([
                r['team'], r['assignee'], r['label'], r['project'],
                r['programme'], r['mapping_name'],
                f"{r['forecast_days']:.2f}", f"{r['actual_days']:.2f}",
                f"{r['delta']:.2f}", r['status'],
            ])
        s = data.get('summary', {})
        w.writerow([])
        w.writerow(['', '', '', '', '', 'TOTAL',
                    f"{s.get('forecast_total_days', 0):.2f}",
                    f"{s.get('actual_total_days', 0):.2f}",
                    f"{s.get('delta', 0):.2f}", ''])
        return buf.getvalue().encode('utf-8-sig')

    @staticmethod
    def export_xlsx(data: dict, sprint_name: str = '') -> bytes:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill

        rows = data.get('rows', [])
        s = data.get('summary', {})

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Sprint FA'

        H_FILL = PatternFill('solid', fgColor='1F4E79')
        STATUS_FILLS = {
            'added':     PatternFill('solid', fgColor='D1FAE5'),
            'removed':   PatternFill('solid', fgColor='FEE2E2'),
            'changed':   PatternFill('solid', fgColor='FEF3C7'),
            'unchanged': PatternFill('solid', fgColor='F9FAFB'),
        }
        CENTER = Alignment(horizontal='center')

        headers = [
            'Team', 'Engineer', 'Label', 'Project', 'Programme', 'Finance Type',
            'Forecast Days', 'Actual Days', 'Delta', 'Status',
        ]
        for ci, h in enumerate(headers, 1):
            c = ws.cell(1, ci, h)
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = H_FILL
            c.alignment = CENTER

        for ri, r in enumerate(rows, 2):
            fill = STATUS_FILLS.get(r['status'], STATUS_FILLS['unchanged'])
            vals = [
                r['team'], r['assignee'], r['label'], r['project'],
                r['programme'], r['mapping_name'],
                float(r['forecast_days']), float(r['actual_days']),
                float(r['delta']), r['status'],
            ]
            for ci, v in enumerate(vals, 1):
                c = ws.cell(ri, ci, v)
                c.fill = fill
                if ci >= 7:
                    c.alignment = CENTER

        # Totals row
        tr = len(rows) + 2
        ws.cell(tr, 6, 'TOTAL').font = Font(bold=True)
        for ci, v in [(7, s.get('forecast_total_days', 0)),
                      (8, s.get('actual_total_days', 0)),
                      (9, s.get('delta', 0))]:
            c = ws.cell(tr, ci, float(v))
            c.font = Font(bold=True)
            c.alignment = CENTER

        col_widths = [18, 20, 16, 20, 18, 18, 14, 14, 10, 12]
        for ci, w in enumerate(col_widths, 1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(ci)].width = w

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()


# ---------------------------------------------------------------------------
# KPI Report — Estimate % Accuracy
# ---------------------------------------------------------------------------

_MONTH_NAMES = [
    'January', 'February', 'March', 'April', 'May', 'June',
    'July', 'August', 'September', 'October', 'November', 'December',
]

# (label, bg_hex, text_hex) — ordered best→worst
_KPI_BANDS = [
    ('> 90%',     '#dcfce7', '#15803d'),
    ('> 80%',     '#ecfccb', '#3f6212'),
    ('In Range',  '#fef3c7', '#92400e'),
    ('> 70%',     '#ffedd5', '#9a3412'),
    ('> 60%',     '#fee2e2', '#991b1b'),
    ('> 50%',     '#fecaca', '#7f1d1d'),
    ('< 50%',     '#fca5a5', '#450a0a'),
    ('Exception', '#ede9fe', '#4c1d95'),
]

KPI_BAND_BG    = {b[0]: b[1] for b in _KPI_BANDS}
KPI_BAND_TEXT  = {b[0]: b[2] for b in _KPI_BANDS}
KPI_BAND_ORDER = [b[0] for b in _KPI_BANDS]


def _kpi_accuracy_band(total_cost, estimate_value, estimate_with_contingency, comment):
    """Return the accuracy band label for a single project row."""
    total      = float(total_cost or 0)
    estimate   = float(estimate_value or 0)
    estimate_wc = float(estimate_with_contingency or 0)

    if comment and comment.strip():
        return 'Exception'

    if estimate <= 0:
        return '—'

    # Over estimate but within contingency → "In Range"
    if total >= estimate and total <= estimate_wc:
        return 'In Range'

    accuracy = total / estimate * 100
    if accuracy >= 90:
        return '> 90%'
    if accuracy >= 80:
        return '> 80%'
    if accuracy >= 70:
        return '> 70%'
    if accuracy >= 60:
        return '> 60%'
    if accuracy >= 50:
        return '> 50%'
    return '< 50%'


class KPIReportService:

    SLUG = 'kpi-estimate-accuracy'

    @staticmethod
    def get_months():
        """Return distinct year-months where completed projects exist, newest first."""
        from apps.projects.models import Project

        rows = (
            Project.objects.filter(
                status=Project.STATUS_COMPLETED,
                completed_sprint__isnull=False,
            )
            .values('completed_sprint__end_date__year', 'completed_sprint__end_date__month')
            .distinct()
        )

        months = sorted(
            {(r['completed_sprint__end_date__year'], r['completed_sprint__end_date__month'])
             for r in rows if r['completed_sprint__end_date__year']},
            reverse=True,
        )

        return [
            {
                'value': f'{y:04d}-{m:02d}',
                'label': f'{_MONTH_NAMES[m - 1]} {y}',
            }
            for y, m in months
        ]

    @staticmethod
    def get_data(month_str):
        """Return full KPI report data for the given month (YYYY-MM)."""
        from apps.projects.models import Project
        from apps.projects.utils import get_tshirt_size
        from apps.sprint_forecast.models import ProjectActuals

        from .models import KPIReportComment

        try:
            year  = int(month_str[:4])
            month = int(month_str[5:7])
        except (ValueError, IndexError, AttributeError):
            return None

        projects = list(
            Project.objects.filter(
                status=Project.STATUS_COMPLETED,
                completed_sprint__isnull=False,
                completed_sprint__end_date__year=year,
                completed_sprint__end_date__month=month,
            )
            .select_related('programme', 'assigned_team', 'completed_sprint')
            .prefetch_related('project_collaborators__team')
            .order_by('name')
        )

        project_ids = [p.id for p in projects]

        actuals_map = {
            a.project_id: a
            for a in ProjectActuals.objects.filter(project_id__in=project_ids)
        }

        comment_map = {
            c.project_id: c.comment
            for c in KPIReportComment.objects.filter(
                project_id__in=project_ids, month=month_str
            )
        }

        rows = []
        for project in projects:
            actuals        = actuals_map.get(project.id)
            estimate_value = float(actuals.estimate_value) if actuals else 0.0
            estimate_wc    = float(actuals.estimate_value_with_contingency) if actuals else 0.0
            total_cost     = float(actuals.total_cost_till_date) if actuals else 0.0

            comment  = comment_map.get(project.id, '')
            tshirt   = get_tshirt_size(estimate_value) if estimate_value > 0 else '—'
            band     = _kpi_accuracy_band(total_cost, estimate_value, estimate_wc, comment)

            accuracy_pct = (
                round(total_cost / estimate_value * 100, 2)
                if estimate_value > 0 else None
            )

            collaborators = [pc.team.name for pc in project.project_collaborators.all()]

            rows.append({
                'project_id':                     project.id,
                'project':                        project.name,
                'programme':                      project.programme.name if project.programme else '—',
                'assigned_team':                  project.assigned_team.name if project.assigned_team else '—',
                'collaborators':                  collaborators,
                'collaborators_str':              ', '.join(collaborators),
                'estimate_value':                 estimate_value,
                'estimate_value_with_contingency': estimate_wc,
                'total_cost_till_date':           total_cost,
                'tshirt_size':                    tshirt,
                'accuracy_pct':                   accuracy_pct,
                'accuracy_band':                  band,
                'comment':                        comment,
                'band_bg':                        KPI_BAND_BG.get(band, '#f9fafb'),
                'band_text':                      KPI_BAND_TEXT.get(band, '#6b7280'),
            })

        def _chart_bands(sizes):
            counts: dict = {}
            for r in rows:
                if r['tshirt_size'] in sizes:
                    b = r['accuracy_band']
                    counts[b] = counts.get(b, 0) + 1
            return counts

        return {
            'month':       month_str,
            'month_label': f'{_MONTH_NAMES[month - 1]} {year}',
            'rows':        rows,
            'charts': {
                'xs_s':   _chart_bands({'XS', 'S'}),
                'm_plus': _chart_bands({'M', 'L', 'XL'}),
            },
        }

    @staticmethod
    def get_configure_data(month_str):
        """Return projects for the month with their existing comments."""
        from apps.projects.models import Project

        from .models import KPIReportComment

        try:
            year  = int(month_str[:4])
            month = int(month_str[5:7])
        except (ValueError, IndexError, AttributeError):
            return None

        projects = list(
            Project.objects.filter(
                status=Project.STATUS_COMPLETED,
                completed_sprint__isnull=False,
                completed_sprint__end_date__year=year,
                completed_sprint__end_date__month=month,
            )
            .select_related('programme')
            .order_by('name')
        )

        project_ids = [p.id for p in projects]
        comment_map = {
            c.project_id: c
            for c in KPIReportComment.objects.filter(
                project_id__in=project_ids, month=month_str
            )
        }

        result = []
        for p in projects:
            obj = comment_map.get(p.id)
            result.append({
                'project_id':   p.id,
                'project_name': p.name,
                'programme':    p.programme.name if p.programme else '—',
                'comment_id':   obj.id if obj else None,
                'comment':      obj.comment if obj else '',
            })

        return {
            'month':    month_str,
            'projects': result,
        }

    @staticmethod
    def save_comments(month_str, comments, user=None):
        """Upsert/delete comments for the given month."""
        from .models import KPIReportComment

        for item in comments:
            project_id = item['project_id']
            text = item.get('comment', '').strip()
            if text:
                KPIReportComment.objects.update_or_create(
                    project_id=project_id,
                    month=month_str,
                    defaults={'comment': text, 'updated_by': user},
                )
            else:
                KPIReportComment.objects.filter(
                    project_id=project_id, month=month_str
                ).delete()

    # ── Exports ───────────────────────────────────────────────────────────────

    @staticmethod
    def export_csv(data: dict) -> bytes:
        rows        = data.get('rows', [])
        month_label = data.get('month_label', '')

        buf = io.StringIO()
        w   = csv.writer(buf)
        w.writerow([
            'Programme', 'Project', 'Assigned Team', 'Collaborators',
            'Estimate Value (£)', 'Estimate Value With Contingency (£)',
            'Total Cost Till Date (£)', 'T-Shirt Size',
            'Estimate % Accuracy', 'Accuracy Band', 'Comment',
        ])
        for r in rows:
            acc = f"{r['accuracy_pct']:.2f}%" if r['accuracy_pct'] is not None else '—'
            w.writerow([
                r['programme'], r['project'], r['assigned_team'],
                r['collaborators_str'],
                f"{r['estimate_value']:,.2f}",
                f"{r['estimate_value_with_contingency']:,.2f}",
                f"{r['total_cost_till_date']:,.2f}",
                r['tshirt_size'], acc,
                r['accuracy_band'], r['comment'],
            ])
        return buf.getvalue().encode('utf-8-sig')

    @staticmethod
    def export_xlsx(data: dict) -> bytes:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        rows        = data.get('rows', [])
        month_label = data.get('month_label', '')

        BAND_STYLES = {
            '> 90%':     ('dcfce7', '15803d'),
            '> 80%':     ('ecfccb', '3f6212'),
            'In Range':  ('fef3c7', '92400e'),
            '> 70%':     ('ffedd5', '9a3412'),
            '> 60%':     ('fee2e2', '991b1b'),
            '> 50%':     ('fecaca', '7f1d1d'),
            '< 50%':     ('fca5a5', '450a0a'),
            'Exception': ('ede9fe', '4c1d95'),
            '—':         ('f9fafb', '6b7280'),
        }

        H_FILL = PatternFill('solid', fgColor='1F4E79')
        CENTER  = Alignment(horizontal='center')
        GBP_FMT = '£#,##0.00'

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'KPI Estimate Accuracy'

        # Title row
        ws.merge_cells('A1:K1')
        title_cell = ws.cell(1, 1, f'KPI Report — Estimate % Accuracy | {month_label}')
        title_cell.font = Font(bold=True, size=13)

        headers = [
            'Programme', 'Project', 'Assigned Team', 'Collaborators',
            'Estimate Value', 'Est. w/Contingency', 'Total Cost Till Date',
            'T-Shirt', '% Accuracy', 'Accuracy Band', 'Comment',
        ]
        for ci, h in enumerate(headers, 1):
            c = ws.cell(3, ci, h)
            c.font   = Font(bold=True, color='FFFFFF')
            c.fill   = H_FILL
            c.alignment = CENTER

        for ri, r in enumerate(rows, 4):
            band        = r['accuracy_band']
            bg, fg      = BAND_STYLES.get(band, BAND_STYLES['—'])
            band_fill   = PatternFill('solid', fgColor=bg)
            band_font   = Font(color=fg, bold=True)

            row_vals = [
                r['programme'], r['project'], r['assigned_team'],
                r['collaborators_str'],
                r['estimate_value'],
                r['estimate_value_with_contingency'],
                r['total_cost_till_date'],
                r['tshirt_size'],
                r['accuracy_pct'],
                r['accuracy_band'],
                r['comment'],
            ]
            for ci, v in enumerate(row_vals, 1):
                c = ws.cell(ri, ci, v)
                if ci in (5, 6, 7):
                    c.number_format = GBP_FMT
                    c.alignment = CENTER
                elif ci == 8:
                    c.alignment = CENTER
                elif ci == 9:
                    if v is not None:
                        c.number_format = '0.00"%"'
                    c.alignment = CENTER
                elif ci == 10:
                    c.fill      = band_fill
                    c.font      = band_font
                    c.alignment = CENTER

        col_widths = [20, 28, 18, 24, 18, 22, 20, 10, 14, 14, 32]
        for ci, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()

    @staticmethod
    def export_pdf(data: dict) -> bytes:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        rows        = data.get('rows', [])
        month_label = data.get('month_label', '')

        def _rl_color(hex_str):
            h = hex_str.lstrip('#')
            return colors.Color(
                int(h[0:2], 16) / 255,
                int(h[2:4], 16) / 255,
                int(h[4:6], 16) / 255,
            )

        BAND_BG_RL = {k: _rl_color(v) for k, v in KPI_BAND_BG.items()}

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=landscape(A4),
            leftMargin=12 * mm, rightMargin=12 * mm,
            topMargin=12 * mm, bottomMargin=12 * mm,
        )

        styles  = getSampleStyleSheet()
        story   = []

        story.append(Paragraph('KPI Report — Estimate % Accuracy', styles['Title']))
        story.append(Paragraph(f'Month: {month_label}', styles['Normal']))
        story.append(Spacer(1, 4 * mm))

        HDR_BG   = colors.Color(31 / 255, 78 / 255, 121 / 255)
        HDR_FONT = colors.white
        ALT_ROW  = colors.Color(0.97, 0.97, 0.97)

        table_data = [[
            'Programme', 'Project', 'Team', 'Collaborators',
            'Estimate (£)', 'Est. w/Ctg (£)', 'Total Cost (£)',
            'T-Shirt', '% Acc.', 'Band', 'Comment',
        ]]

        band_rows = []
        for r in rows:
            acc = f"{r['accuracy_pct']:.1f}%" if r['accuracy_pct'] is not None else '—'
            table_data.append([
                r['programme'], r['project'], r['assigned_team'],
                r['collaborators_str'] or '—',
                f"£{r['estimate_value']:,.0f}",
                f"£{r['estimate_value_with_contingency']:,.0f}",
                f"£{r['total_cost_till_date']:,.0f}",
                r['tshirt_size'], acc,
                r['accuracy_band'],
                r['comment'] or '',
            ])
            band_rows.append(r['accuracy_band'])

        col_w = [25*mm, 30*mm, 22*mm, 25*mm, 20*mm, 22*mm, 20*mm, 12*mm, 14*mm, 14*mm, 30*mm]

        tbl = Table(table_data, colWidths=col_w, repeatRows=1)

        cmds = [
            ('BACKGROUND',   (0, 0), (-1, 0),  HDR_BG),
            ('TEXTCOLOR',    (0, 0), (-1, 0),  HDR_FONT),
            ('FONTNAME',     (0, 0), (-1, 0),  'Helvetica-Bold'),
            ('FONTSIZE',     (0, 0), (-1, -1), 7),
            ('ALIGN',        (4, 0), (9, -1),  'RIGHT'),
            ('ALIGN',        (9, 0), (9, -1),  'CENTER'),
            ('GRID',         (0, 0), (-1, -1), 0.3, colors.lightgrey),
            ('VALIGN',       (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING',   (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING',(0, 0), (-1, -1), 3),
        ]

        for i, band in enumerate(band_rows):
            row_idx = i + 1
            if i % 2 == 1:
                cmds.append(('ROWBACKGROUNDS', (0, row_idx), (-1, row_idx), [ALT_ROW]))
            bg = BAND_BG_RL.get(band)
            if bg:
                cmds.append(('BACKGROUND', (9, row_idx), (9, row_idx), bg))

        tbl.setStyle(TableStyle(cmds))
        story.append(tbl)

        doc.build(story)
        buf.seek(0)
        return buf.read()


# ---------------------------------------------------------------------------
# Monthly Finance Report — Aggregated Recharge Pivot by Project & Programme
# ---------------------------------------------------------------------------

class MonthlyFinanceReportService:

    SLUG = 'monthly-finance'

    @staticmethod
    def get_months():
        """Return distinct year-months where sprints exist, newest first."""
        from apps.sprints.models import Sprint

        rows = (
            Sprint.objects
            .filter(end_date__isnull=False)
            .values('end_date__year', 'end_date__month')
            .distinct()
            .order_by('-end_date__year', '-end_date__month')
        )

        return [
            {
                'value': f"{r['end_date__year']:04d}-{r['end_date__month']:02d}",
                'label': f"{_MONTH_NAMES[r['end_date__month'] - 1]} {r['end_date__year']}",
            }
            for r in rows
            if r['end_date__year'] and r['end_date__month']
        ]

    @staticmethod
    def get_data(month_str):
        from decimal import Decimal

        from apps.sprint_forecast.models import (
            IMPORT_TYPE_ACTUAL,
            RECHARGE_TYPE_ACTUAL,
            ProjectActuals,
            Recharge,
            SprintImportReviewComplete,
        )
        from apps.sprints.models import Sprint

        try:
            year  = int(month_str[:4])
            month = int(month_str[5:7])
        except (ValueError, IndexError, AttributeError):
            return None

        sprints = list(
            Sprint.objects.filter(
                end_date__year=year,
                end_date__month=month,
            ).order_by('start_date')
        )

        month_label = f"{_MONTH_NAMES[month - 1]} {year}"

        if not sprints:
            return {
                'month': month_str,
                'month_label': month_label,
                'error': 'no_sprints',
                'rows': [],
                'sprints': [],
                'sprint_info': [],
            }

        sprint_ids = [s.id for s in sprints]

        sprints_with_actuals = set(
            SprintImportReviewComplete.objects.filter(
                sprint_id__in=sprint_ids,
                import_type=IMPORT_TYPE_ACTUAL,
            ).values_list('sprint_id', flat=True)
        )

        sprint_info = [
            {
                'id': s.id,
                'name': s.sprint_name,
                'has_actuals': s.id in sprints_with_actuals,
            }
            for s in sprints
        ]

        total_sprints   = len(sprint_ids)
        actuals_count   = len(sprints_with_actuals)

        if actuals_count < total_sprints:
            return {
                'month': month_str,
                'month_label': month_label,
                'error': 'incomplete_actuals',
                'total_sprints': total_sprints,
                'actuals_sprints': actuals_count,
                'missing_sprints': total_sprints - actuals_count,
                'sprint_info': sprint_info,
                'rows': [],
            }

        recharges = list(
            Recharge.objects.filter(
                sprint_id__in=sprint_ids,
                type=RECHARGE_TYPE_ACTUAL,
            ).select_related('project', 'programme')
        )

        project_ids = {r.project_id for r in recharges if r.project_id}
        actuals_code_map = {
            a.project_id: a.code
            for a in ProjectActuals.objects.filter(project_id__in=project_ids)
        }

        pivot = {}
        for r in recharges:
            key = (r.project_id, r.programme_id)
            if key not in pivot:
                proj = r.project
                prog = r.programme
                pivot[key] = {
                    'project_id':   r.project_id,
                    'project_code': actuals_code_map.get(r.project_id, '') or '',
                    'project_name': proj.name if proj else '—',
                    'programme_id': r.programme_id,
                    'programme':    prog.name if prog else '—',
                    'total_cost':   Decimal('0'),
                    'total_days':   Decimal('0'),
                }
            pivot[key]['total_cost'] += r.total_cost
            pivot[key]['total_days'] += r.total_days

        rows = sorted(
            pivot.values(),
            key=lambda x: (x['project_code'] or x['project_name'], x['programme']),
        )

        for row in rows:
            row['total_cost'] = _r2(row['total_cost'])
            row['total_days'] = _r2(row['total_days'])

        grand_total_cost = sum(r['total_cost'] for r in rows)
        grand_total_days = sum(r['total_days'] for r in rows)

        return {
            'month':       month_str,
            'month_label': month_label,
            'sprints':     [{'id': s.id, 'name': s.sprint_name} for s in sprints],
            'sprint_info': sprint_info,
            'rows':        rows,
            'summary': {
                'total_projects':    len({r['project_id'] for r in rows}),
                'total_rows':        len(rows),
                'grand_total_cost':  _r2(grand_total_cost),
                'grand_total_days':  _r2(grand_total_days),
            },
        }

    @staticmethod
    def export_csv(data: dict) -> bytes:
        rows        = data.get('rows', [])
        month_label = data.get('month_label', '')
        s           = data.get('summary', {})

        buf = io.StringIO()
        w   = csv.writer(buf)
        w.writerow(['Monthly Finance Report', month_label])
        w.writerow([])
        w.writerow(['Project Code', 'Project Name', 'Programme', 'Total Days', 'Total Cost (£)'])
        for r in rows:
            w.writerow([
                r['project_code'] or '—',
                r['project_name'],
                r['programme'],
                f"{r['total_days']:.2f}",
                f"{r['total_cost']:,.2f}",
            ])
        w.writerow([])
        w.writerow([
            '', '', 'TOTAL',
            f"{s.get('grand_total_days', 0):.2f}",
            f"{s.get('grand_total_cost', 0):,.2f}",
        ])
        return buf.getvalue().encode('utf-8-sig')

    @staticmethod
    def export_xlsx(data: dict) -> bytes:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        rows        = data.get('rows', [])
        month_label = data.get('month_label', '')
        s           = data.get('summary', {})

        H_FILL   = PatternFill('solid', fgColor='1F4E79')
        ALT_FILL = PatternFill('solid', fgColor='F0F4F8')
        TOT_FILL = PatternFill('solid', fgColor='DBEAFE')
        CENTER   = Alignment(horizontal='center')
        GBP_FMT  = '£#,##0.00'

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Monthly Finance'

        ws.merge_cells('A1:E1')
        tc = ws.cell(1, 1, f'Monthly Finance Report | {month_label}')
        tc.font = Font(bold=True, size=13)

        headers = ['Project Code', 'Project Name', 'Programme', 'Total Days', 'Total Cost (£)']
        for ci, h in enumerate(headers, 1):
            c = ws.cell(3, ci, h)
            c.font      = Font(bold=True, color='FFFFFF')
            c.fill      = H_FILL
            c.alignment = CENTER

        for ri, r in enumerate(rows, 4):
            fill = ALT_FILL if ri % 2 == 0 else None
            vals = [
                r['project_code'] or '—',
                r['project_name'],
                r['programme'],
                float(r['total_days']),
                float(r['total_cost']),
            ]
            for ci, v in enumerate(vals, 1):
                c = ws.cell(ri, ci, v)
                if fill:
                    c.fill = fill
                if ci == 4:
                    c.alignment = CENTER
                elif ci == 5:
                    c.number_format = GBP_FMT
                    c.alignment     = CENTER

        tr = len(rows) + 4
        ws.cell(tr, 3, 'TOTAL').font = Font(bold=True)
        ws.cell(tr, 3).fill          = TOT_FILL
        for ci, v in [
            (4, float(s.get('grand_total_days', 0))),
            (5, float(s.get('grand_total_cost', 0))),
        ]:
            c               = ws.cell(tr, ci, v)
            c.font          = Font(bold=True)
            c.fill          = TOT_FILL
            c.alignment     = CENTER
            if ci == 5:
                c.number_format = GBP_FMT

        col_widths = [18, 32, 24, 14, 18]
        for ci, cw in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = cw

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()


# ---------------------------------------------------------------------------
# Registry
# ---------------------------------------------------------------------------

REPORT_REGISTRY = {
    "demand-capacity": DemandCapacityService,
}
