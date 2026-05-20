"""
Custom Report Query Engine — translates a report config into Django ORM queries
and returns structured data suitable for the chosen visualization.
"""
import io
import csv
from importlib import import_module
from typing import Any, Dict, List

from django.core.exceptions import ValidationError
from django.db.models import Count, Sum, Avg, Min, Max, F, Q

from .data_sources import get_data_source

AGG_COUNT          = 'count'
AGG_COUNT_DISTINCT = 'count_distinct'
AGG_SUM            = 'sum'
AGG_AVERAGE        = 'average'
AGG_MINIMUM        = 'minimum'
AGG_MAXIMUM        = 'maximum'

_AGG_FNS = {
    AGG_COUNT:          lambda f: Count(f),
    AGG_COUNT_DISTINCT: lambda f: Count(f, distinct=True),
    AGG_SUM:            lambda f: Sum(f),
    AGG_AVERAGE:        lambda f: Avg(f),
    AGG_MINIMUM:        lambda f: Min(f),
    AGG_MAXIMUM:        lambda f: Max(f),
}

_FILTER_OPS = {
    'eq':          lambda f, v: Q(**{f: v}),
    'neq':         lambda f, v: ~Q(**{f: v}),
    'gt':          lambda f, v: Q(**{f'{f}__gt': v}),
    'gte':         lambda f, v: Q(**{f'{f}__gte': v}),
    'lt':          lambda f, v: Q(**{f'{f}__lt': v}),
    'lte':         lambda f, v: Q(**{f'{f}__lte': v}),
    'contains':    lambda f, v: Q(**{f'{f}__icontains': v}),
    'starts_with': lambda f, v: Q(**{f'{f}__istartswith': v}),
    'is_null':     lambda f, v: Q(**{f'{f}__isnull': bool(v)}),
    'in':          lambda f, v: Q(**{f'{f}__in': v if isinstance(v, list) else [v]}),
}

MAX_ROWS = 5000


def _import_model(model_path: str):
    module_path, cls_name = model_path.rsplit('.', 1)
    return getattr(import_module(module_path), cls_name)


def _ann_key(field_key: str) -> str:
    """'team__name'  →  'team_name'  (safe annotation alias)."""
    return field_key.replace('__', '_')


def _annotate_fk_fields(qs, field_keys: List[str]):
    """For any field that uses __ traversal, add an F() annotation so values() can use it."""
    ann = {_ann_key(fk): F(fk) for fk in field_keys if '__' in fk}
    if ann:
        qs = qs.annotate(**ann)
    return qs, {fk: _ann_key(fk) for fk in field_keys if '__' in fk}


def _resolve_key(field_key: str, fk_map: Dict) -> str:
    return fk_map.get(field_key, field_key)


def _build_filter_q(filters: List[Dict]) -> Q:
    q = Q()
    for f in (filters or []):
        field = f.get('field', '')
        op    = f.get('operator', 'eq')
        value = f.get('value')
        if not field or op not in _FILTER_OPS:
            continue
        try:
            q &= _FILTER_OPS[op](field, value)
        except Exception:
            continue
    return q


def _value_annotations(values_cfg: List[Dict]) -> Dict[str, Any]:
    ann = {}
    for i, v in enumerate(values_cfg or []):
        field = v.get('field') or 'id'
        agg   = v.get('aggregation') or AGG_COUNT
        fn    = _AGG_FNS.get(agg, _AGG_FNS[AGG_COUNT])
        ann[f'_val{i}'] = fn(field)
    return ann


def _value_label(v: Dict, field_map: Dict) -> str:
    agg   = (v.get('aggregation') or AGG_COUNT).replace('_', ' ').upper()
    field = v.get('field') or 'id'
    return f'{agg}({field_map.get(field, field)})'


# ── Public entry point ────────────────────────────────────────────────────────

def execute(config: Dict, data_source_key: str) -> Dict:
    """
    Execute a report config and return structured data.

    config keys:
      visualization: str
      fields:        [str, ...]   (table raw columns)
      filters:       [{field, operator, value}, ...]
      values:        [{field, aggregation, series_type?}, ...]
      axis:          str   (bar / line / pie x-axis)
      legend:        str   (bar / line grouping / color)
      rows:          str   (pivot / heatmap row dimension)
      columns:       str   (pivot / heatmap column dimension)
    """
    ds = get_data_source(data_source_key)
    if not ds:
        raise ValidationError(f'Unknown data source: {data_source_key!r}')

    model = _import_model(ds['model'])
    qs    = model.objects.all()

    # Apply data-source-level base filters (e.g. import_type=FORECAST)
    base_filt = _build_filter_q(ds.get('base_filters', []))
    if base_filt:
        qs = qs.filter(base_filt)

    filt = _build_filter_q(config.get('filters', []))
    if filt:
        qs = qs.filter(filt)

    field_map = {f['key']: f['label'] for f in ds['fields']}
    viz       = config.get('visualization', 'table')

    if viz == 'table':
        return _table(qs, config, field_map, ds)
    if viz in ('pivot', 'heatmap'):
        result = _pivot(qs, config, field_map)
        result['type'] = viz
        return result
    if viz == 'pie':
        return _pie(qs, config, field_map)
    if viz == 'card':
        return _card(qs, config, field_map)
    # bar, stacked_bar, column, stacked_column, line, combo
    return _chart(qs, config, field_map, viz)


# ── Visualization builders ────────────────────────────────────────────────────

def _table(qs, config, field_map, ds) -> Dict:
    all_field_keys = [f['key'] for f in ds['fields']]
    fields     = config.get('fields') or all_field_keys
    values_cfg = config.get('values') or []

    # Grouped table: group by selected fields and compute aggregations
    if fields and values_cfg:
        qs, fk_map   = _annotate_fk_fields(qs, fields)
        group_keys   = [_resolve_key(fk, fk_map) for fk in fields]
        ann          = _value_annotations(values_cfg)
        rows         = list(qs.values(*group_keys).annotate(**ann).order_by(*group_keys)[:MAX_ROWS])
        dim_cols     = [{'key': _resolve_key(fk, fk_map), 'label': field_map.get(fk, fk)} for fk in fields]
        val_cols     = [{'key': f'_val{i}', 'label': _value_label(v, field_map)} for i, v in enumerate(values_cfg)]
        return {'type': 'table', 'columns': dim_cols + val_cols, 'rows': rows, 'total': len(rows)}

    # Raw table
    qs, fk_map = _annotate_fk_fields(qs, fields)
    val_keys   = [_resolve_key(fk, fk_map) for fk in fields]
    rows       = list(qs.values(*val_keys)[:MAX_ROWS])
    columns    = [{'key': _resolve_key(fk, fk_map), 'label': field_map.get(fk, fk)} for fk in fields]
    return {'type': 'table', 'columns': columns, 'rows': rows, 'total': len(rows)}


def _chart(qs, config, field_map, viz) -> Dict:
    axis          = config.get('axis') or ''
    legend        = config.get('legend') or ''
    values_cfg    = config.get('values') or [{'field': 'id', 'aggregation': AGG_COUNT}]

    if not axis:
        return {'type': viz, 'error': 'axis field is required for this visualization'}

    group_fields = [axis] + ([legend] if legend and legend != axis else [])
    qs, fk_map   = _annotate_fk_fields(qs, group_fields)
    axis_key     = _resolve_key(axis,   fk_map)
    legend_key   = _resolve_key(legend, fk_map) if legend else ''
    group_keys   = [axis_key] + ([legend_key] if legend_key and legend_key != axis_key else [])

    ann   = _value_annotations(values_cfg)
    rows  = list(qs.values(*group_keys).annotate(**ann).order_by(*group_keys))

    v_labels = [_value_label(v, field_map) for v in values_cfg]

    if legend_key and len(group_keys) > 1:
        all_x = sorted({str(r.get(axis_key, '')) for r in rows})
        all_l = sorted({str(r.get(legend_key, '')) for r in rows})
        datasets = [
            {
                'label':       lv,
                'series_type': values_cfg[0].get('series_type', 'bar') if values_cfg else 'bar',
                'data':        [next((r.get('_val0', 0) for r in rows
                                     if str(r.get(axis_key, '')) == x
                                     and str(r.get(legend_key, '')) == lv), 0)
                                for x in all_x],
            }
            for lv in all_l
        ]
        return {'type': viz, 'labels': all_x, 'datasets': datasets,
                'axis_label': field_map.get(axis, axis), 'legend_label': field_map.get(legend, legend)}
    else:
        labels   = [str(r.get(axis_key, '')) for r in rows]
        datasets = [
            {
                'label':       v_labels[i] if i < len(v_labels) else 'Value',
                'series_type': v.get('series_type', 'bar'),
                'data':        [r.get(f'_val{i}', 0) for r in rows],
            }
            for i, v in enumerate(values_cfg)
        ]
        return {'type': viz, 'labels': labels, 'datasets': datasets,
                'axis_label': field_map.get(axis, axis)}


def _pie(qs, config, field_map) -> Dict:
    axis       = config.get('axis') or 'id'
    values_cfg = config.get('values') or [{'field': 'id', 'aggregation': AGG_COUNT}]
    ann        = _value_annotations(values_cfg)

    qs, fk_map = _annotate_fk_fields(qs, [axis])
    axis_key   = _resolve_key(axis, fk_map)

    rows   = list(qs.values(axis_key).annotate(**ann).order_by(axis_key))
    labels = [str(r.get(axis_key, '')) for r in rows]
    data   = [r.get('_val0', 0) for r in rows]

    return {'type': 'pie', 'labels': labels,
            'datasets': [{'label': field_map.get(axis, axis), 'data': data}]}


def _card(qs, config, field_map) -> Dict:
    values_cfg = config.get('values') or [{'field': 'id', 'aggregation': AGG_COUNT}]
    cards = []
    for i, v in enumerate(values_cfg[:6]):
        fn  = _AGG_FNS.get(v.get('aggregation') or AGG_COUNT, _AGG_FNS[AGG_COUNT])
        ann = {f'_val{i}': fn(v.get('field') or 'id')}
        try:
            row = qs.aggregate(**ann)
            val = row.get(f'_val{i}', 0)
        except Exception:
            val = None
        cards.append({'label': _value_label(v, field_map), 'value': val})
    return {'type': 'card', 'cards': cards}


def _pivot(qs, config, field_map) -> Dict:
    rows_f     = config.get('rows')    or ''
    cols_f     = config.get('columns') or ''
    values_cfg = config.get('values')  or [{'field': 'id', 'aggregation': AGG_COUNT}]

    if not rows_f or not cols_f:
        return {'type': 'pivot', 'error': 'rows and columns fields are required for pivot'}

    ann = _value_annotations(values_cfg)
    qs, fk_map = _annotate_fk_fields(qs, [rows_f, cols_f])
    rows_key   = _resolve_key(rows_f, fk_map)
    cols_key   = _resolve_key(cols_f, fk_map)

    data = list(qs.values(rows_key, cols_key).annotate(**ann).order_by(rows_key, cols_key))

    all_rows = sorted({str(r[rows_key]) for r in data})
    all_cols = sorted({str(r[cols_key]) for r in data})

    cell = {rv: {cv: 0 for cv in all_cols} for rv in all_rows}
    for d in data:
        rv  = str(d[rows_key])
        cv  = str(d[cols_key])
        val = d.get('_val0') or 0
        cell[rv][cv] = val

    matrix     = {rv: [cell[rv][cv] for cv in all_cols] for rv in all_rows}
    row_totals = {rv: sum(cell[rv].values()) for rv in all_rows}
    col_totals = {cv: sum(cell[rv][cv] for rv in all_rows) for cv in all_cols}

    return {
        'type':         'pivot',
        'rows':         all_rows,
        'columns':      all_cols,
        'matrix':       matrix,
        'row_totals':   row_totals,
        'col_totals':   list(col_totals.values()),
        'grand_total':  sum(row_totals.values()),
        'rows_label':   field_map.get(rows_f,  rows_f),
        'columns_label':field_map.get(cols_f, cols_f),
    }


# ── Export helpers ────────────────────────────────────────────────────────────

def export_csv(result: Dict) -> str:
    """Serialize table or pivot result to CSV string."""
    buf = io.StringIO()
    w   = csv.writer(buf)

    if result['type'] == 'table':
        cols = result.get('columns', [])
        w.writerow([c['label'] for c in cols])
        for row in result.get('rows', []):
            w.writerow([row.get(c['key'], '') for c in cols])

    elif result['type'] in ('pivot', 'heatmap'):
        cols   = result.get('columns', [])
        header = [result.get('rows_label', 'Row')] + cols + ['Total']
        w.writerow(header)
        matrix   = result.get('matrix', {})
        row_tots = result.get('row_totals', {})
        for rv in result.get('rows', []):
            w.writerow([rv] + list(matrix.get(rv, [])) + [row_tots.get(rv, '')])
        w.writerow(['Total'] + result.get('col_totals', []) + [result.get('grand_total', '')])

    return buf.getvalue()


def export_xlsx(result: Dict) -> bytes:
    """Serialize table or pivot result to XLSX bytes."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active

    header_fill = PatternFill('solid', fgColor='1E3A5F')
    header_font = Font(bold=True, color='FFFFFF')

    def _hdr(cell):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    if result['type'] == 'table':
        cols = result.get('columns', [])
        for ci, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=ci, value=c['label'])
            _hdr(cell)
        for ri, row in enumerate(result.get('rows', []), 2):
            for ci, c in enumerate(cols, 1):
                ws.cell(row=ri, column=ci, value=row.get(c['key'], ''))

    elif result['type'] in ('pivot', 'heatmap'):
        columns = result.get('columns', [])
        ws.cell(row=1, column=1, value=result.get('rows_label', 'Row'))
        _hdr(ws.cell(row=1, column=1))
        for ci, cv in enumerate(columns, 2):
            _hdr(ws.cell(row=1, column=ci, value=cv))
        total_col = len(columns) + 2
        _hdr(ws.cell(row=1, column=total_col, value='Total'))

        matrix   = result.get('matrix', {})
        row_tots = result.get('row_totals', {})
        for ri, rv in enumerate(result.get('rows', []), 2):
            ws.cell(row=ri, column=1, value=rv)
            for ci, val in enumerate(matrix.get(rv, []), 2):
                ws.cell(row=ri, column=ci, value=val)
            ws.cell(row=ri, column=total_col, value=row_tots.get(rv, ''))
        last_row = len(result.get('rows', [])) + 2
        ws.cell(row=last_row, column=1, value='Total')
        for ci, val in enumerate(result.get('col_totals', []), 2):
            ws.cell(row=last_row, column=ci, value=val)
        ws.cell(row=last_row, column=total_col, value=result.get('grand_total', ''))

    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
