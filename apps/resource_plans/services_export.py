# Export service module — imported by api_views for the export endpoint.

from .services import (
    AllocationSetService,
    CapacityService,
    TeamUtilisationService,
)


class ExportService:
    """Build an openpyxl workbook from live plan data with five sheets."""

    _styles_ready = False

    @classmethod
    def _init_styles(cls):
        if cls._styles_ready:
            return
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        thin = Side(style='thin', color='AAAAAA')
        cls._BORDER    = Border(left=thin, right=thin, top=thin, bottom=thin)
        cls._HDR_FILL  = PatternFill('solid', fgColor='4361EE')
        cls._HDR_FONT  = Font(bold=True, color='FFFFFF', size=9)
        cls._TEAM_FILL = PatternFill('solid', fgColor='E9ECEF')
        cls._OVER_FILL = PatternFill('solid', fgColor='F8D7DA')
        cls._WARN_FILL = PatternFill('solid', fgColor='FFF3CD')
        cls._GOOD_FILL = PatternFill('solid', fgColor='D4EDDA')
        cls._CENT      = Alignment(horizontal='center', vertical='center')
        cls._styles_ready = True

    @classmethod
    def _hdr(cls, ws, row, col, value, width=None):
        from openpyxl.styles import Alignment
        cell = ws.cell(row=row, column=col, value=value)
        cell.fill   = cls._HDR_FILL
        cell.font   = cls._HDR_FONT
        cell.border = cls._BORDER
        cell.alignment = cls._CENT
        if width:
            ws.column_dimensions[cell.column_letter].width = width
        return cell

    @classmethod
    def _cell(cls, ws, row, col, value, fill=None, bold=False,
              num_fmt=None, align_right=False):
        from openpyxl.styles import Font, Alignment
        cell = ws.cell(row=row, column=col, value=value)
        cell.border = cls._BORDER
        cell.font   = Font(bold=bold, size=9)
        cell.alignment = Alignment(
            horizontal='right' if align_right else 'left',
            vertical='center',
        )
        if fill:
            cell.fill = fill
        if num_fmt:
            cell.number_format = num_fmt
        return cell

    # ── Sheet 1: Allocation Grid ──────────────────────────────────────────────

    @classmethod
    def _sheet_allocation_grid(cls, wb, version, allocation_set_id, team_id):
        ws = wb.create_sheet('Allocation Grid')
        grid    = AllocationSetService.get_allocations_grid(version, allocation_set_id, team_id)
        sprints = grid.get('sprints', [])
        rows    = grid.get('rows', [])

        cls._hdr(ws, 1, 1, 'Programme', 18)
        cls._hdr(ws, 1, 2, 'Project',   24)
        cls._hdr(ws, 1, 3, 'Team',      16)
        cls._hdr(ws, 1, 4, 'Member',    20)
        cls._hdr(ws, 1, 5, 'Phase',     18)
        cls._hdr(ws, 1, 6, 'Type',      12)
        for si, s in enumerate(sprints):
            cls._hdr(ws, 1, 7 + si, s['name'], 8)

        r = 2
        for row in rows:
            cells = row.get('cells', [])
            fill  = cls._TEAM_FILL if row.get('is_team_row') else None
            bold  = bool(row.get('is_team_row'))
            cls._cell(ws, r, 1, row.get('programme_name', ''), fill=fill, bold=bold)
            cls._cell(ws, r, 2, row.get('project_name',   ''), fill=fill, bold=bold)
            cls._cell(ws, r, 3, row.get('team_name',      ''), fill=fill, bold=bold)
            cls._cell(ws, r, 4, row.get('member_name',    ''), fill=fill, bold=bold)
            cls._cell(ws, r, 5, row.get('phase_name',     ''), fill=fill)
            cls._cell(ws, r, 6, row.get('assignment_type',''), fill=fill)
            for si, c in enumerate(cells[:len(sprints)]):
                val  = float(c.get('effective_days', 0) or 0)
                cfil = cls._OVER_FILL if c.get('is_over') else fill
                cls._cell(ws, r, 7 + si,
                          val if val else None,
                          fill=cfil, num_fmt='0.00', align_right=True)
            r += 1

        ws.freeze_panes = 'G2'
        ws.row_dimensions[1].height = 28
        return ws

    # ── Sheet 2: Net Capacity ─────────────────────────────────────────────────

    @classmethod
    def _sheet_net_capacity(cls, wb, version, team_id):
        ws = wb.create_sheet('Net Capacity')
        cap     = CapacityService.get_capacity_grid(version, team_id)
        sprints = cap.get('sprints', [])
        rows    = cap.get('rows', [])

        cls._hdr(ws, 1, 1, 'Team',   16)
        cls._hdr(ws, 1, 2, 'Member', 22)
        for si, s in enumerate(sprints):
            cls._hdr(ws, 1, 3 + si, s['name'], 8)

        r = 2
        for row in rows:
            cls._cell(ws, r, 1, row.get('team_name', ''))
            cls._cell(ws, r, 2, row.get('member_name', ''))
            for si, c in enumerate(row.get('cells', [])[:len(sprints)]):
                val = c.get('net_capacity')
                cls._cell(ws, r, 3 + si,
                          float(val) if val is not None else None,
                          num_fmt='0.00', align_right=True)
            r += 1

        ws.freeze_panes = 'C2'
        return ws

    # ── Sheet 3: Holidays & Leaves ────────────────────────────────────────────

    @classmethod
    def _sheet_holidays_leaves(cls, wb, version, team_id):
        ws = wb.create_sheet('Holidays & Leaves')
        cap     = CapacityService.get_absences_grid(version, team_id)
        sprints = cap.get('sprints', [])
        rows    = cap.get('rows', [])

        cls._hdr(ws, 1, 1, 'Team',         16)
        cls._hdr(ws, 1, 2, 'Member',        22)
        cls._hdr(ws, 1, 3, 'Sprint',        12)
        cls._hdr(ws, 1, 4, 'Working Days',  14)
        cls._hdr(ws, 1, 5, 'Holiday Days',  14)
        cls._hdr(ws, 1, 6, 'Leave Days',    14)
        cls._hdr(ws, 1, 7, 'Net Capacity',  14)

        r = 2
        for row in rows:
            for si, c in enumerate(row.get('cells', [])[:len(sprints)]):
                sname = sprints[si]['name'] if si < len(sprints) else ''
                cls._cell(ws, r, 1, row.get('team_name', ''))
                cls._cell(ws, r, 2, row.get('member_name', ''))
                cls._cell(ws, r, 3, sname)
                cls._cell(ws, r, 4, float(c['working_days']) if c.get('working_days') else None,
                          num_fmt='0.00', align_right=True)
                cls._cell(ws, r, 5, float(c['holiday_days']) if c.get('holiday_days') else None,
                          num_fmt='0.00', align_right=True)
                cls._cell(ws, r, 6, float(c['leave_days'])   if c.get('leave_days')   else None,
                          num_fmt='0.00', align_right=True)
                cls._cell(ws, r, 7, float(c['net_capacity'])  if c.get('net_capacity')  else None,
                          num_fmt='0.00', align_right=True)
                r += 1

        ws.freeze_panes = 'D2'
        return ws

    # ── Sheet 4: Utilisation Summary ──────────────────────────────────────────

    @classmethod
    def _sheet_utilisation(cls, wb, version, allocation_set_id, team_id):
        ws = wb.create_sheet('Utilisation Summary')
        team_ids = [team_id] if team_id else None
        data    = TeamUtilisationService.get_team_utilisation(
            version, allocation_set_id, team_ids=team_ids
        )
        sprints = data.get('sprints', [])
        rows    = data.get('rows', [])

        cls._hdr(ws, 1, 1, 'Team',            20)
        cls._hdr(ws, 1, 2, 'Sprint',           12)
        cls._hdr(ws, 1, 3, 'Net Capacity (d)', 18)
        cls._hdr(ws, 1, 4, 'Allocated (d)',    16)
        cls._hdr(ws, 1, 5, 'Utilisation %',    14)
        cls._hdr(ws, 1, 6, 'Over-allocated',    14)

        r = 2
        for row in rows:
            for si, c in enumerate(row.get('cells', [])[:len(sprints)]):
                sname = sprints[si]['name'] if si < len(sprints) else ''
                util  = c.get('utilisation_pct')
                is_over = c.get('is_over')
                fill  = (cls._OVER_FILL if is_over else
                         cls._GOOD_FILL if util is not None and util >= 85 else
                         cls._WARN_FILL if util is not None and util >= 50 else None)
                cls._cell(ws, r, 1, row.get('team_name', ''))
                cls._cell(ws, r, 2, sname)
                cls._cell(ws, r, 3, c.get('net_capacity'),
                          num_fmt='0.00', align_right=True)
                cls._cell(ws, r, 4, c.get('allocated_days'),
                          fill=fill, num_fmt='0.00', align_right=True)
                cls._cell(ws, r, 5, util,
                          fill=fill, num_fmt='0.0', align_right=True)
                cls._cell(ws, r, 6, 'Yes' if is_over else '',
                          fill=fill if is_over else None)
                r += 1

        ws.freeze_panes = 'C2'
        return ws

    # ── Sheet 5: Conflicts ────────────────────────────────────────────────────

    @classmethod
    def _sheet_conflicts(cls, wb, version):
        from .models import Conflict
        ws = wb.create_sheet('Conflicts')

        cls._hdr(ws, 1, 1, 'Type',        20)
        cls._hdr(ws, 1, 2, 'Severity',    12)
        cls._hdr(ws, 1, 3, 'Status',      12)
        cls._hdr(ws, 1, 4, 'Sprint',      14)
        cls._hdr(ws, 1, 5, 'Description', 60)

        conflicts = (
            Conflict.objects
            .filter(allocation_set__version=version)
            .select_related('affected_sprint')
            .order_by('severity', 'conflict_type')
        )

        r = 2
        for c in conflicts:
            sev_fill = cls._OVER_FILL if c.severity == 'ERROR' else cls._WARN_FILL
            cls._cell(ws, r, 1, c.conflict_type)
            cls._cell(ws, r, 2, c.severity, fill=sev_fill)
            cls._cell(ws, r, 3, c.status)
            cls._cell(ws, r, 4,
                      c.affected_sprint.sprint_name if c.affected_sprint else '')
            cls._cell(ws, r, 5, c.description)
            r += 1

        if r == 2:
            ws.cell(row=2, column=1, value='No conflicts recorded for this version.')

        return ws

    # ── Public entry point ────────────────────────────────────────────────────

    @classmethod
    def build_workbook(cls, version, allocation_set_id=None, team_id=None):
        """Return an openpyxl Workbook with five sheets."""
        import openpyxl
        cls._init_styles()
        wb = openpyxl.Workbook()
        wb.remove(wb.active)   # drop the default empty sheet

        cls._sheet_allocation_grid(wb, version, allocation_set_id, team_id)
        cls._sheet_net_capacity(wb, version, team_id)
        cls._sheet_holidays_leaves(wb, version, team_id)
        cls._sheet_utilisation(wb, version, allocation_set_id, team_id)
        cls._sheet_conflicts(wb, version)

        return wb
