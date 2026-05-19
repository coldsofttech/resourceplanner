import io
import logging
from collections import defaultdict
from datetime import datetime, timezone

from django.core.exceptions import ValidationError
from django.core.paginator import Paginator

from .models import (
    Win, WinEntry,
    CATEGORY_CHOICES, CATEGORY_DELIVERY, CATEGORY_OPERATIONAL,
    MonthlyWin, MonthlyWinSurvey, MonthlyWinSurveyNomination, MonthlyWinResult,
    TeamProductOwner,
)

logger = logging.getLogger(__name__)


# ── Weekly Win Service ────────────────────────────────────────────────────────

class WinService:

    @staticmethod
    def next_week_number() -> int:
        from apps.configurations.services import ConfigurationService
        start = ConfigurationService.get_int('WIN_START_NUMBER', 1)
        latest = Win.objects.order_by('-week_number').values_list('week_number', flat=True).first()
        return max(start, (latest + 1) if latest is not None else start)

    @staticmethod
    def create_win(week_start_date, user=None) -> Win:
        import datetime
        week_number = WinService.next_week_number()
        win = Win(
            week_start_date=week_start_date,
            week_end_date=week_start_date + datetime.timedelta(days=6),
            week_number=week_number,
            created_by=user,
        )
        win.full_clean()
        win.save()
        return win

    @staticmethod
    def list_wins(page=1, page_size=20):
        qs = Win.objects.all()
        paginator = Paginator(qs, page_size)
        p = paginator.get_page(page)
        return {
            'results': list(p.object_list),
            'total_count': paginator.count,
            'total_pages': paginator.num_pages,
            'current_page': p.number,
            'page_size': page_size,
            'has_next': p.has_next(),
            'has_previous': p.has_previous(),
        }

    @staticmethod
    def get_win(pk) -> Win:
        return Win.objects.get(pk=pk)

    @staticmethod
    def add_entry(win_pk, team_id, title, description='', user=None) -> WinEntry:
        win = Win.objects.get(pk=win_pk)
        if not title or not title.strip():
            raise ValidationError({'title': 'Title is required.'})
        entry = WinEntry(
            win=win,
            team_id=team_id,
            title=title.strip(),
            description=(description or '').strip(),
            created_by=user,
        )
        entry.full_clean()
        entry.save()
        return entry

    @staticmethod
    def update_entry(entry_pk, title, description=None) -> WinEntry:
        entry = WinEntry.objects.get(pk=entry_pk)
        if not title or not title.strip():
            raise ValidationError({'title': 'Title is required.'})
        entry.title = title.strip()
        if description is not None:
            entry.description = description.strip()
        entry.full_clean()
        entry.save()
        return entry

    @staticmethod
    def delete_entry(entry_pk):
        WinEntry.objects.filter(pk=entry_pk).delete()

    @staticmethod
    def review_complete(win_pk, user=None):
        """Mark a win as reviewed, generate Word doc, email to configured recipients."""
        win = Win.objects.prefetch_related('entries__team').get(pk=win_pk)
        if win.status == Win.STATUS_REVIEW_COMPLETE:
            raise ValidationError('This week has already been marked as reviewed.')

        doc_bytes = WinService._build_docx(win)
        WinService._send_review_email(win, doc_bytes)

        win.status = Win.STATUS_REVIEW_COMPLETE
        win.reviewed_at = datetime.now(tz=timezone.utc)
        win.reviewed_by = user
        win.save(update_fields=['status', 'reviewed_at', 'reviewed_by'])
        return win

    @staticmethod
    def _build_docx(win: Win) -> bytes:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Title
        title = doc.add_heading('', level=0)
        run = title.add_run('Weekly Wins Review')
        run.font.color.rgb = RGBColor(0x1e, 0x40, 0xaf)

        # Week heading
        week_label = (
            f'Week {win.week_number} — '
            f'{win.week_start_date.strftime("%d %b %Y")} to {win.week_end_date.strftime("%d %b %Y")}'
        )
        doc.add_heading(week_label, level=1)

        # Group entries by team
        by_team = defaultdict(list)
        for entry in win.entries.all():
            by_team[entry.team.name].append(entry)

        if not by_team:
            doc.add_paragraph('No wins recorded for this week.')
        else:
            for team_name in sorted(by_team.keys()):
                entries = by_team[team_name]
                doc.add_heading(team_name, level=2)
                for entry in entries:
                    para = doc.add_paragraph(style='List Bullet')
                    run = para.add_run(entry.title)
                    run.bold = True
                    if entry.description:
                        para.add_run(f'\n{entry.description}')

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    @staticmethod
    def _send_review_email(win: Win, doc_bytes: bytes):
        from django.core.mail import EmailMessage
        from apps.configurations.services import ConfigurationService

        recipients_raw = ConfigurationService.get_str('WINS_REVIEW_EMAIL_RECIPIENTS', '')
        recipients = [r.strip() for r in recipients_raw.split(',') if r.strip()]
        if not recipients:
            logger.warning('WINS_REVIEW_EMAIL_RECIPIENTS is not configured — skipping email.')
            return

        subject = (
            f'Weekly Wins — Week {win.week_number} '
            f'({win.week_start_date.strftime("%d %b %Y")} – {win.week_end_date.strftime("%d %b %Y")})'
        )
        body = (
            f'Please find attached the Weekly Wins review document for '
            f'Week {win.week_number} '
            f'({win.week_start_date.strftime("%d %b %Y")} – {win.week_end_date.strftime("%d %b %Y")}).\n\n'
            f'This is an automated message from the Resource Planner.'
        )
        filename = f'weekly_wins_week_{win.week_number}.docx'

        email = EmailMessage(subject=subject, body=body, to=recipients)
        email.attach(filename, doc_bytes, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        email.send(fail_silently=True)

    @staticmethod
    def get_report_data(week_ids=None, date_from=None, date_to=None):
        """Return rows + summary for the report, filtered by week IDs or date range."""
        qs = Win.objects.prefetch_related('entries__team')

        if week_ids:
            qs = qs.filter(pk__in=week_ids)
        elif date_from and date_to:
            qs = qs.filter(week_start_date__gte=date_from, week_end_date__lte=date_to)
        elif date_from:
            qs = qs.filter(week_start_date__gte=date_from)
        elif date_to:
            qs = qs.filter(week_end_date__lte=date_to)

        qs = qs.order_by('week_number')

        rows = []
        team_counts = defaultdict(lambda: {'team_name': '', 'win_count': 0})

        for win in qs:
            for entry in win.entries.all():
                rows.append({
                    'entry_id': entry.id,
                    'team_id': entry.team_id,
                    'team_name': entry.team.name,
                    'week_number': win.week_number,
                    'week_start_date': str(win.week_start_date),
                    'week_end_date': str(win.week_end_date),
                    'title': entry.title,
                    'description': entry.description,
                })
                team_counts[entry.team_id]['team_name'] = entry.team.name
                team_counts[entry.team_id]['win_count'] += 1

        summary = [
            {'team_id': tid, 'team_name': v['team_name'], 'win_count': v['win_count']}
            for tid, v in sorted(team_counts.items(), key=lambda x: x[1]['team_name'])
        ]

        return {'rows': rows, 'summary': summary}


# ── Monthly Win Service ───────────────────────────────────────────────────────

class MonthlyWinService:

    @staticmethod
    def list_monthly_wins(page=1, page_size=20):
        qs = MonthlyWin.objects.all()
        paginator = Paginator(qs, page_size)
        p = paginator.get_page(page)
        return {
            'results': list(p.object_list),
            'total_count': paginator.count,
            'total_pages': paginator.num_pages,
            'current_page': p.number,
            'page_size': page_size,
            'has_next': p.has_next(),
            'has_previous': p.has_previous(),
        }

    @staticmethod
    def create_monthly_win(name, win_ids, phase1_deadline=None, user=None) -> MonthlyWin:
        if not name or not name.strip():
            raise ValidationError({'name': 'Name is required.'})
        mw = MonthlyWin(name=name.strip(), created_by=user, phase1_deadline=phase1_deadline)
        mw.save()
        if win_ids:
            mw.wins.set(Win.objects.filter(pk__in=win_ids))
        return mw

    @staticmethod
    def update_monthly_win(pk, name=None, win_ids=None, phase1_deadline=None, phase2_deadline=None) -> MonthlyWin:
        mw = MonthlyWin.objects.get(pk=pk)
        if mw.status != MonthlyWin.STATUS_DRAFT:
            raise ValidationError('Only draft monthly wins can be edited.')
        if name is not None:
            mw.name = name.strip()
        if phase1_deadline is not None:
            mw.phase1_deadline = phase1_deadline
        if phase2_deadline is not None:
            mw.phase2_deadline = phase2_deadline
        mw.save()
        if win_ids is not None:
            mw.wins.set(Win.objects.filter(pk__in=win_ids))
        return mw

    @staticmethod
    def get_monthly_win(pk) -> MonthlyWin:
        return MonthlyWin.objects.prefetch_related(
            'wins', 'surveys__teams', 'surveys__nominations__entry__team', 'results__entry__team',
        ).get(pk=pk)

    @staticmethod
    def launch_phase1(pk, user=None) -> MonthlyWin:
        mw = MonthlyWin.objects.prefetch_related('wins').get(pk=pk)
        if mw.status != MonthlyWin.STATUS_DRAFT:
            raise ValidationError('Phase 1 can only be launched from Draft status.')

        wins = list(mw.wins.all())
        if not wins:
            raise ValidationError('Select at least one week before launching Phase 1.')

        # Find all teams that have entries in the selected weeks
        team_ids = (
            WinEntry.objects
            .filter(win__in=wins)
            .values_list('team_id', flat=True)
            .distinct()
        )

        # Find POs for those teams
        po_mappings = (
            TeamProductOwner.objects
            .filter(team_id__in=team_ids, is_active=True)
            .select_related('user', 'team')
        )

        if not po_mappings.exists():
            raise ValidationError('No Product Owners are configured for the teams with wins in the selected weeks.')

        # Group by PO: each PO gets one survey covering all their teams
        po_teams = defaultdict(list)
        for mapping in po_mappings:
            po_teams[mapping.user_id].append(mapping.team)

        po_users = {m.user_id: m.user for m in po_mappings}

        for user_id, teams in po_teams.items():
            survey = MonthlyWinSurvey.objects.create(
                monthly_win=mw,
                phase=MonthlyWinSurvey.PHASE_1,
                recipient_id=user_id,
            )
            survey.teams.set(teams)

        mw.status = MonthlyWin.STATUS_PHASE1_OPEN
        mw.save(update_fields=['status'])

        # Send emails + in-app notifications
        MonthlyWinService._send_phase1_emails(mw)
        MonthlyWinService._notify_phase1_pos(mw)

        return mw

    @staticmethod
    def _build_survey_entry_labels(mw: MonthlyWin):
        """
        For each entry in the selected weeks, build its survey label:
        'Week N: [i] Title'
        where i is the index within (team, week) group.
        """
        wins = list(mw.wins.order_by('week_number'))
        entries = (
            WinEntry.objects
            .filter(win__in=wins)
            .select_related('win', 'team')
            .order_by('team__name', 'win__week_number', 'created_at')
        )

        counters = defaultdict(int)  # (team_id, win_id) → count
        labels = {}
        for entry in entries:
            key = (entry.team_id, entry.win_id)
            counters[key] += 1
            idx = counters[key]
            suffix = f' [{idx}]' if counters[key] > 1 else ''
            labels[entry.id] = f'Week {entry.win.week_number}: {suffix}{entry.title}'.strip()

        return labels, list(entries)

    @staticmethod
    def _notify_phase1_pos(mw: MonthlyWin):
        from apps.notifications.models import Notification
        from apps.notifications.services import NotificationService
        from apps.configurations.services import ConfigurationService
        site_url = ConfigurationService.get_str('SITE_URL', '').rstrip('/')
        for survey in mw.surveys.filter(phase=MonthlyWinSurvey.PHASE_1, status=MonthlyWinSurvey.STATUS_PENDING).select_related('recipient'):
            link = f'{site_url}/wins/survey/{survey.token}/'
            deadline_str = f' Deadline: {mw.phase1_deadline.strftime("%d %b %Y %H:%M")}.' if mw.phase1_deadline else ''
            NotificationService.create(
                user=survey.recipient,
                title=f'Monthly Wins — {mw.name}: Phase 1 Voting Open',
                body=f'You have been invited to vote for Monthly Wins.{deadline_str}',
                link=link,
                notification_type=Notification.TYPE_MONTHLY_WINS_PHASE1,
            )

    @staticmethod
    def _notify_phase2_pos(mw: MonthlyWin):
        from apps.notifications.models import Notification
        from apps.notifications.services import NotificationService
        from apps.configurations.services import ConfigurationService
        site_url = ConfigurationService.get_str('SITE_URL', '').rstrip('/')
        for survey in mw.surveys.filter(phase=MonthlyWinSurvey.PHASE_2, status=MonthlyWinSurvey.STATUS_PENDING).select_related('recipient'):
            link = f'{site_url}/wins/survey/{survey.token}/'
            deadline_str = f' Deadline: {mw.phase2_deadline.strftime("%d %b %Y %H:%M")}.' if mw.phase2_deadline else ''
            NotificationService.create(
                user=survey.recipient,
                title=f'Monthly Wins — {mw.name}: Phase 2 Final Voting Open',
                body=f'Phase 2 voting is now open. Please cast your final votes.{deadline_str}',
                link=link,
                notification_type=Notification.TYPE_MONTHLY_WINS_PHASE2,
            )

    @staticmethod
    def _send_phase1_emails(mw: MonthlyWin):
        from apps.configurations.services import ConfigurationService
        site_url = ConfigurationService.get_str('SITE_URL', '').rstrip('/')

        labels, _ = MonthlyWinService._build_survey_entry_labels(mw)

        for survey in mw.surveys.filter(phase=MonthlyWinSurvey.PHASE_1, status=MonthlyWinSurvey.STATUS_PENDING):
            MonthlyWinService._send_survey_email(survey, site_url, labels)

    @staticmethod
    def _send_survey_email(survey: MonthlyWinSurvey, site_url: str, entry_labels: dict = None):
        from django.core.mail import send_mail
        from apps.configurations.services import ConfigurationService

        if site_url is None:
            site_url = ConfigurationService.get_str('SITE_URL', '').rstrip('/')

        link = f'{site_url}/wins/survey/{survey.token}/'
        deadline_str = ''
        if survey.monthly_win.phase1_deadline and survey.phase == MonthlyWinSurvey.PHASE_1:
            deadline_str = f'\n\nPlease respond by {survey.monthly_win.phase1_deadline.strftime("%d %b %Y %H:%M")}.'
        elif survey.monthly_win.phase2_deadline and survey.phase == MonthlyWinSurvey.PHASE_2:
            deadline_str = f'\n\nPlease respond by {survey.monthly_win.phase2_deadline.strftime("%d %b %Y %H:%M")}.'

        phase_label = 'Phase 1' if survey.phase == MonthlyWinSurvey.PHASE_1 else 'Phase 2 (Final Selection)'
        subject = f'Monthly Wins — {survey.monthly_win.name} — {phase_label} Survey'

        if survey.phase == MonthlyWinSurvey.PHASE_1:
            team_list = ', '.join(t.name for t in survey.teams.all())
            body = (
                f'Hi {survey.recipient.get_full_name() or survey.recipient.email},\n\n'
                f'You have been asked to vote for the best wins from your team(s): {team_list}.\n\n'
                f'Please click the link below to complete the survey:\n{link}{deadline_str}\n\n'
                f'This is an automated message from the Resource Planner.'
            )
        else:
            body = (
                f'Hi {survey.recipient.get_full_name() or survey.recipient.email},\n\n'
                f'Phase 2 of the Monthly Wins is now open. '
                f'Please select the best wins across all teams.\n\n'
                f'Please click the link below to complete the survey:\n{link}{deadline_str}\n\n'
                f'This is an automated message from the Resource Planner.'
            )

        try:
            send_mail(subject, body, None, [survey.recipient.email], fail_silently=True)
            survey.sent_at = datetime.now(tz=timezone.utc)
            survey.save(update_fields=['sent_at'])
        except Exception:
            logger.exception('Failed to send survey email for survey %s', survey.pk)

    @staticmethod
    def send_reminder(survey_pk) -> MonthlyWinSurvey:
        from apps.configurations.services import ConfigurationService
        survey = MonthlyWinSurvey.objects.select_related('monthly_win', 'recipient').get(pk=survey_pk)
        if survey.status != MonthlyWinSurvey.STATUS_PENDING:
            raise ValidationError('Reminders can only be sent for pending surveys.')

        site_url = ConfigurationService.get_str('SITE_URL', '').rstrip('/')
        MonthlyWinService._send_survey_email(survey, site_url)

        survey.reminder_count += 1
        survey.last_reminder_at = datetime.now(tz=timezone.utc)
        survey.save(update_fields=['reminder_count', 'last_reminder_at'])
        return survey

    @staticmethod
    def override_survey(survey_pk) -> MonthlyWinSurvey:
        survey = MonthlyWinSurvey.objects.get(pk=survey_pk)
        if survey.status == MonthlyWinSurvey.STATUS_COMPLETED:
            raise ValidationError('Survey is already completed.')
        survey.status = MonthlyWinSurvey.STATUS_OVERRIDDEN
        survey.completed_at = datetime.now(tz=timezone.utc)
        survey.save(update_fields=['status', 'completed_at'])
        return survey

    @staticmethod
    def dismiss_nomination(nomination_pk, reason='') -> MonthlyWinSurveyNomination:
        nom = MonthlyWinSurveyNomination.objects.get(pk=nomination_pk)
        nom.is_dismissed = True
        nom.dismissed_reason = (reason or '').strip()
        nom.save(update_fields=['is_dismissed', 'dismissed_reason'])
        return nom

    @staticmethod
    def undismiss_nomination(nomination_pk) -> MonthlyWinSurveyNomination:
        nom = MonthlyWinSurveyNomination.objects.get(pk=nomination_pk)
        nom.is_dismissed = False
        nom.dismissed_reason = ''
        nom.save(update_fields=['is_dismissed', 'dismissed_reason'])
        return nom

    @staticmethod
    def complete_phase1(pk) -> MonthlyWin:
        mw = MonthlyWin.objects.get(pk=pk)
        if mw.status != MonthlyWin.STATUS_PHASE1_OPEN:
            raise ValidationError('Phase 1 is not currently open.')
        mw.status = MonthlyWin.STATUS_PHASE1_COMPLETE
        mw.save(update_fields=['status'])
        return mw

    @staticmethod
    def launch_phase2(pk) -> MonthlyWin:
        mw = MonthlyWin.objects.prefetch_related('wins').get(pk=pk)
        if mw.status != MonthlyWin.STATUS_PHASE1_COMPLETE:
            raise ValidationError('Phase 2 can only be launched after Phase 1 is complete.')

        # All active POs in the system receive a phase 2 survey (one per unique user)
        seen_user_ids = set()
        all_pos = TeamProductOwner.objects.filter(is_active=True).select_related('user')
        if not all_pos.exists():
            raise ValidationError('No active Product Owners found.')

        for mapping in all_pos:
            if mapping.user_id in seen_user_ids:
                continue
            seen_user_ids.add(mapping.user_id)
            MonthlyWinSurvey.objects.create(
                monthly_win=mw,
                phase=MonthlyWinSurvey.PHASE_2,
                recipient=mapping.user,
            )

        mw.status = MonthlyWin.STATUS_PHASE2_OPEN
        mw.save(update_fields=['status'])

        MonthlyWinService._send_phase2_emails(mw)
        MonthlyWinService._notify_phase2_pos(mw)

        return mw

    @staticmethod
    def _send_phase2_emails(mw: MonthlyWin):
        from apps.configurations.services import ConfigurationService
        site_url = ConfigurationService.get_str('SITE_URL', '').rstrip('/')
        for survey in mw.surveys.filter(phase=MonthlyWinSurvey.PHASE_2, status=MonthlyWinSurvey.STATUS_PENDING):
            MonthlyWinService._send_survey_email(survey, site_url)

    @staticmethod
    def get_survey_data(token):
        """Return survey + entry options for the survey page (token-based)."""
        survey = (
            MonthlyWinSurvey.objects
            .select_related('monthly_win', 'recipient')
            .prefetch_related('teams', 'monthly_win__wins', 'nominations')
            .get(token=token)
        )
        mw = survey.monthly_win

        if survey.phase == MonthlyWinSurvey.PHASE_1:
            # Entries for the PO's teams within selected weeks
            team_ids = list(survey.teams.values_list('id', flat=True))
            entries = (
                WinEntry.objects
                .filter(win__in=mw.wins.all(), team_id__in=team_ids)
                .select_related('win', 'team')
                .order_by('team__name', 'win__week_number', 'created_at')
            )
        else:
            # Phase 2: entries that were nominated (not dismissed) in phase 1
            nominated_entry_ids = (
                MonthlyWinSurveyNomination.objects
                .filter(survey__monthly_win=mw, survey__phase=MonthlyWinSurvey.PHASE_1, is_dismissed=False)
                .values_list('entry_id', flat=True)
                .distinct()
            )
            entries = (
                WinEntry.objects
                .filter(pk__in=nominated_entry_ids)
                .select_related('win', 'team')
                .order_by('team__name', 'win__week_number', 'created_at')
            )

        # Build labels with indexed numbering per (team, week)
        # First pass: determine total count per (team, win) key
        entries_list = list(entries)
        key_totals = defaultdict(int)
        for entry in entries_list:
            key_totals[(entry.team_id, entry.win_id)] += 1

        counters = defaultdict(int)
        entry_data = []
        for entry in entries_list:
            key = (entry.team_id, entry.win_id)
            counters[key] += 1
            idx = counters[key]
            suffix = f' [{idx}]' if key_totals[key] > 1 else ''
            label = f'Week {entry.win.week_number}:{suffix} {entry.title}'.strip()
            entry_data.append({
                'id': entry.id,
                'label': label,
                'team_id': entry.team_id,
                'team_name': entry.team.name,
                'week_number': entry.win.week_number,
                'title': entry.title,
                'description': entry.description,
            })

        existing_nominations = {
            (n.entry_id, n.category): n.id
            for n in survey.nominations.all()
        }

        return {
            'survey': survey,
            'monthly_win': mw,
            'entries': entry_data,
            'existing_nominations': existing_nominations,
            'categories': [{'value': v, 'label': l} for v, l in CATEGORY_CHOICES],
        }

    @staticmethod
    def submit_survey(token, nominations_data):
        """
        nominations_data: list of {'entry_id': int, 'category': str}
        Phase 1 constraint: max 2 per (team, category)
        Phase 2 constraint: max 2 per category total
        """
        survey = MonthlyWinSurvey.objects.get(token=token)
        if survey.status != MonthlyWinSurvey.STATUS_PENDING:
            raise ValidationError('This survey has already been completed or overridden.')

        # Validate selections
        if survey.phase == MonthlyWinSurvey.PHASE_1:
            # Max 2 per (team, category)
            team_cat_counts = defaultdict(int)
            team_ids = set(survey.teams.values_list('id', flat=True))
            for nom in nominations_data:
                entry = WinEntry.objects.select_related('team').get(pk=nom['entry_id'])
                if entry.team_id not in team_ids:
                    raise ValidationError(f'Entry {nom["entry_id"]} does not belong to your teams.')
                key = (entry.team_id, nom['category'])
                team_cat_counts[key] += 1
                if team_cat_counts[key] > 2:
                    raise ValidationError(
                        f'You may select at most 2 {nom["category"]} wins per team.'
                    )
        else:
            # Phase 2: max 2 per category total
            cat_counts = defaultdict(int)
            for nom in nominations_data:
                cat_counts[nom['category']] += 1
                if cat_counts[nom['category']] > 2:
                    raise ValidationError(f'You may select at most 2 {nom["category"]} wins.')

        # Cross-category validation: same entry cannot appear in both categories
        entry_categories = defaultdict(set)
        for nom in nominations_data:
            entry_categories[nom['entry_id']].add(nom['category'])
        for entry_id, cats in entry_categories.items():
            if len(cats) > 1:
                raise ValidationError('A win cannot be selected for both Delivery and Operational Excellence.')

        # Clear existing nominations and save new ones
        survey.nominations.all().delete()
        for nom in nominations_data:
            MonthlyWinSurveyNomination.objects.create(
                survey=survey,
                entry_id=nom['entry_id'],
                category=nom['category'],
            )

        survey.status = MonthlyWinSurvey.STATUS_COMPLETED
        survey.completed_at = datetime.now(tz=timezone.utc)
        survey.save(update_fields=['status', 'completed_at'])
        return survey

    @staticmethod
    def declare_winners(pk) -> MonthlyWin:
        mw = MonthlyWin.objects.get(pk=pk)
        if mw.status != MonthlyWin.STATUS_PHASE2_OPEN:
            raise ValidationError('Winners can only be declared while Phase 2 is open.')

        # Count votes per (entry, category) from non-dismissed phase 2 nominations
        vote_counts = defaultdict(int)
        nominations = (
            MonthlyWinSurveyNomination.objects
            .filter(survey__monthly_win=mw, survey__phase=MonthlyWinSurvey.PHASE_2, is_dismissed=False)
            .values('entry_id', 'category')
        )
        for nom in nominations:
            vote_counts[(nom['entry_id'], nom['category'])] += 1

        # Top 2 per category
        MonthlyWinResult.objects.filter(monthly_win=mw).delete()
        for category, _ in CATEGORY_CHOICES:
            cat_votes = [
                (entry_id, count)
                for (entry_id, cat), count in vote_counts.items()
                if cat == category
            ]
            cat_votes.sort(key=lambda x: -x[1])
            for rank, (entry_id, count) in enumerate(cat_votes[:2], start=1):
                MonthlyWinResult.objects.create(
                    monthly_win=mw,
                    entry_id=entry_id,
                    category=category,
                    rank=rank,
                    vote_count=count,
                )

        mw.status = MonthlyWin.STATUS_DECLARED
        mw.save(update_fields=['status'])
        return mw

    @staticmethod
    def get_teams_for_preview(mw_pk):
        """Return teams that have entries in the selected weeks (for Phase 1 preview dropdown)."""
        mw = MonthlyWin.objects.prefetch_related('wins').get(pk=mw_pk)
        from apps.delivery_teams.models import DeliveryTeam
        team_ids = (
            WinEntry.objects
            .filter(win__in=mw.wins.all())
            .values_list('team_id', flat=True)
            .distinct()
        )
        return list(DeliveryTeam.objects.filter(pk__in=team_ids).order_by('name').values('id', 'name'))

    @staticmethod
    def get_preview_survey_data(mw_pk, phase, team_id=None):
        """
        Simulate survey data for admin preview (no real survey object needed).
        Phase 1: shows entries for the specified team from selected weeks.
        Phase 2: shows Phase 1 nominated (non-dismissed) entries.
        """
        mw = MonthlyWin.objects.prefetch_related('wins').get(pk=mw_pk)

        if phase == MonthlyWinSurvey.PHASE_1:
            if not team_id:
                raise ValidationError('team_id is required for Phase 1 preview.')
            entries = (
                WinEntry.objects
                .filter(win__in=mw.wins.all(), team_id=team_id)
                .select_related('win', 'team')
                .order_by('win__week_number', 'created_at')
            )
        else:
            nominated_entry_ids = (
                MonthlyWinSurveyNomination.objects
                .filter(survey__monthly_win=mw, survey__phase=MonthlyWinSurvey.PHASE_1, is_dismissed=False)
                .values_list('entry_id', flat=True)
                .distinct()
            )
            entries = (
                WinEntry.objects
                .filter(pk__in=nominated_entry_ids)
                .select_related('win', 'team')
                .order_by('team__name', 'win__week_number', 'created_at')
            )

        entries_list = list(entries)
        key_totals = defaultdict(int)
        for entry in entries_list:
            key_totals[(entry.team_id, entry.win_id)] += 1

        counters = defaultdict(int)
        entry_data = []
        for entry in entries_list:
            key = (entry.team_id, entry.win_id)
            counters[key] += 1
            idx = counters[key]
            suffix = f' [{idx}]' if key_totals[key] > 1 else ''
            label = f'Week {entry.win.week_number}:{suffix} {entry.title}'.strip()
            entry_data.append({
                'id': entry.id,
                'label': label,
                'team_id': entry.team_id,
                'team_name': entry.team.name,
                'week_number': entry.win.week_number,
                'title': entry.title,
                'description': entry.description,
            })

        return {
            'phase': phase,
            'entries': entry_data,
            'categories': [{'value': v, 'label': l} for v, l in CATEGORY_CHOICES],
        }

    @staticmethod
    def get_admin_survey_data(survey_pk):
        """Return survey + entry data for admin override form."""
        survey = (
            MonthlyWinSurvey.objects
            .select_related('monthly_win', 'recipient')
            .prefetch_related('teams', 'monthly_win__wins', 'nominations')
            .get(pk=survey_pk)
        )
        mw = survey.monthly_win

        if survey.phase == MonthlyWinSurvey.PHASE_1:
            team_ids = list(survey.teams.values_list('id', flat=True))
            entries = (
                WinEntry.objects
                .filter(win__in=mw.wins.all(), team_id__in=team_ids)
                .select_related('win', 'team')
                .order_by('team__name', 'win__week_number', 'created_at')
            )
        else:
            nominated_ids = (
                MonthlyWinSurveyNomination.objects
                .filter(survey__monthly_win=mw, survey__phase=MonthlyWinSurvey.PHASE_1, is_dismissed=False)
                .values_list('entry_id', flat=True)
                .distinct()
            )
            entries = (
                WinEntry.objects
                .filter(pk__in=nominated_ids)
                .select_related('win', 'team')
                .order_by('team__name', 'win__week_number', 'created_at')
            )

        entries_list = list(entries)
        key_totals = defaultdict(int)
        for entry in entries_list:
            key_totals[(entry.team_id, entry.win_id)] += 1

        counters = defaultdict(int)
        entry_data = []
        for entry in entries_list:
            key = (entry.team_id, entry.win_id)
            counters[key] += 1
            idx = counters[key]
            suffix = f' [{idx}]' if key_totals[key] > 1 else ''
            label = f'Week {entry.win.week_number}:{suffix} {entry.title}'.strip()
            entry_data.append({
                'id': entry.id,
                'label': label,
                'team_id': entry.team_id,
                'team_name': entry.team.name,
                'week_number': entry.win.week_number,
                'title': entry.title,
                'description': entry.description,
            })

        existing_noms = [
            {'entry_id': n.entry_id, 'category': n.category}
            for n in survey.nominations.all()
        ]

        return {
            'phase': survey.phase,
            'recipient_name': survey.recipient.get_full_name() or survey.recipient.email,
            'team_names': [t.name for t in survey.teams.all()],
            'entries': entry_data,
            'categories': [{'value': v, 'label': l} for v, l in CATEGORY_CHOICES],
            'existing_nominations': existing_noms,
        }

    @staticmethod
    def override_survey_with_nominations(survey_pk, nominations_data) -> MonthlyWinSurvey:
        """
        Admin fills in survey on behalf of a PO and marks it as overridden.
        Applies same validation rules as submit_survey.
        """
        survey = MonthlyWinSurvey.objects.select_related('monthly_win').prefetch_related('teams').get(pk=survey_pk)
        if survey.status != MonthlyWinSurvey.STATUS_PENDING:
            raise ValidationError('Only pending surveys can be overridden.')

        if survey.phase == MonthlyWinSurvey.PHASE_1:
            team_cat_counts = defaultdict(int)
            team_ids = set(survey.teams.values_list('id', flat=True))
            for nom in nominations_data:
                entry = WinEntry.objects.select_related('team').get(pk=nom['entry_id'])
                if entry.team_id not in team_ids:
                    raise ValidationError(f'Entry {nom["entry_id"]} does not belong to this survey\'s teams.')
                key = (entry.team_id, nom['category'])
                team_cat_counts[key] += 1
                if team_cat_counts[key] > 2:
                    raise ValidationError(f'You may select at most 2 {nom["category"]} wins per team.')
        else:
            cat_counts = defaultdict(int)
            for nom in nominations_data:
                cat_counts[nom['category']] += 1
                if cat_counts[nom['category']] > 2:
                    raise ValidationError(f'You may select at most 2 {nom["category"]} wins.')

        entry_categories = defaultdict(set)
        for nom in nominations_data:
            entry_categories[nom['entry_id']].add(nom['category'])
        for entry_id, cats in entry_categories.items():
            if len(cats) > 1:
                raise ValidationError('A win cannot be selected for both Delivery and Operational Excellence.')

        survey.nominations.all().delete()
        for nom in nominations_data:
            MonthlyWinSurveyNomination.objects.create(
                survey=survey,
                entry_id=nom['entry_id'],
                category=nom['category'],
            )

        survey.status = MonthlyWinSurvey.STATUS_OVERRIDDEN
        survey.completed_at = datetime.now(tz=timezone.utc)
        survey.save(update_fields=['status', 'completed_at'])
        return survey

    @staticmethod
    def get_phase1_nominations(monthly_win_pk):
        """Return all phase 1 nominations for admin review."""
        return (
            MonthlyWinSurveyNomination.objects
            .filter(survey__monthly_win_id=monthly_win_pk, survey__phase=MonthlyWinSurvey.PHASE_1)
            .select_related('survey__recipient', 'entry__team', 'entry__win')
            .order_by('category', 'entry__team__name', 'entry__win__week_number')
        )


# ── Team Product Owner Service ────────────────────────────────────────────────

class TeamProductOwnerService:

    @staticmethod
    def list_for_team(team_id):
        return TeamProductOwner.objects.filter(team_id=team_id).select_related('user', 'team')

    @staticmethod
    def list_all():
        return TeamProductOwner.objects.all().select_related('user', 'team')

    @staticmethod
    def add(team_id, user_id) -> TeamProductOwner:
        tpo, created = TeamProductOwner.objects.get_or_create(team_id=team_id, user_id=user_id)
        if not created and not tpo.is_active:
            tpo.is_active = True
            tpo.save(update_fields=['is_active'])
        return tpo

    @staticmethod
    def remove(pk):
        TeamProductOwner.objects.filter(pk=pk).delete()
